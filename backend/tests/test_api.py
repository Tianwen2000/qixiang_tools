import base64
import hashlib
import io
import json
import sqlite3
import tempfile
import uuid
import zipfile
from pathlib import Path
from urllib.parse import parse_qs

import qrcode
from docx import Document
from ebooklib import epub
from fastapi.testclient import TestClient
from Crypto.Hash import keccak
from Crypto.PublicKey import RSA
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pillow_heif import from_pillow
from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.main import app
from app.tools.rabbit_cipher import RabbitCipher


client = TestClient(app)


def make_png_bytes(color: tuple[int, int, int] = (255, 0, 0), size: tuple[int, int] = (120, 80)) -> bytes:
    stream = io.BytesIO()
    Image.new("RGB", size, color).save(stream, format="PNG")
    return stream.getvalue()


def make_noise_png_bytes(size: tuple[int, int] = (480, 480)) -> bytes:
    stream = io.BytesIO()
    Image.effect_noise(size, 100).convert("L").convert("RGB").save(stream, format="PNG")
    return stream.getvalue()


def make_jpg_bytes(color: tuple[int, int, int] = (255, 0, 0), size: tuple[int, int] = (120, 80)) -> bytes:
    stream = io.BytesIO()
    Image.new("RGB", size, color).save(stream, format="JPEG", quality=90)
    return stream.getvalue()


def make_gif_bytes() -> bytes:
    first = Image.new("RGBA", (64, 64), (255, 0, 0, 255))
    second = Image.new("RGBA", (64, 64), (0, 0, 255, 255))
    stream = io.BytesIO()
    first.save(stream, format="GIF", save_all=True, append_images=[second], duration=[120, 120], loop=0)
    return stream.getvalue()


def make_zip_bytes(files: dict[str, bytes]) -> bytes:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    stream.seek(0)
    return stream.getvalue()


def make_docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_heading("测试文档", level=0)
    document.add_paragraph(text)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def make_pdf_bytes(text: str) -> bytes:
    stream = io.BytesIO()
    pdf = canvas.Canvas(stream, pagesize=A4)
    pdf.setFont("Helvetica", 14)
    pdf.drawString(72, A4[1] - 72, text)
    pdf.save()
    stream.seek(0)
    return stream.getvalue()


def make_epub_bytes(title: str, text: str) -> bytes:
    book = epub.EpubBook()
    book.set_identifier("test-book")
    book.set_title(title)
    book.set_language("zh-CN")
    chapter = epub.EpubHtml(title="正文", file_name="chapter-1.xhtml", lang="zh-CN")
    chapter.content = f"<h1>{title}</h1><p>{text}</p>"
    book.add_item(chapter)
    book.toc = (chapter,)
    book.spine = ["nav", chapter]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())

    with tempfile.NamedTemporaryFile(suffix=".epub") as temp_file:
        epub.write_epub(temp_file.name, book)
        return Path(temp_file.name).read_bytes()


def make_txt_bytes(text: str) -> bytes:
    return text.encode("utf-8")


def make_heic_bytes(color: tuple[int, int, int] = (90, 140, 210), size: tuple[int, int] = (120, 80)) -> bytes:
    image = Image.new("RGB", size, color)
    stream = io.BytesIO()
    from_pillow(image).save(stream, format="HEIF")
    stream.seek(0)
    return stream.getvalue()


def make_xlsx_bytes(rows: list[list[object]], extra_sheet: bool = False) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "数据表"
    for row in rows:
        worksheet.append(row)
    if extra_sheet:
        second = workbook.create_sheet("第二页")
        second.append(["编号", "值"])
        second.append([1, "A"])
    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream.getvalue()


def make_csv_bytes(rows: list[list[str]]) -> bytes:
    stream = io.StringIO()
    for row in rows:
        stream.write(",".join(row) + "\n")
    return stream.getvalue().encode("utf-8")


def make_pptx_bytes(slides: list[list[str]]) -> bytes:
    presentation = Presentation()
    for index, lines in enumerate(slides):
        slide_layout = presentation.slide_layouts[1]
        slide = presentation.slides.add_slide(slide_layout)
        slide.shapes.title.text = f"第 {index + 1} 页"
        slide.placeholders[1].text = "\n".join(lines)
    if presentation.slides:
        first = presentation.slides[0]
        first.shapes.title.text = "封面"
    stream = io.BytesIO()
    presentation.save(stream)
    stream.seek(0)
    return stream.getvalue()


def test_health() -> None:
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["data"]["status"] == "ok"


def test_meta_endpoints() -> None:
    categories = client.get("/api/categories")
    assert categories.status_code == 200
    category_slugs = [item["slug"] for item in categories.json()["data"]]
    assert "dev" in category_slugs
    assert "ops" in category_slugs
    assert "game" in category_slugs
    assert category_slugs.index("ops") == category_slugs.index("dev") + 1

    tools = client.get("/api/tools", params={"keyword": "json"})
    assert tools.status_code == 200
    assert any(item["slug"] == "json-format" for item in tools.json()["data"])

    ops_tools = client.get("/api/tools", params={"category": "ops"})
    assert ops_tools.status_code == 200
    ops_tool_items = ops_tools.json()["data"]
    ops_tool_slugs = [item["slug"] for item in ops_tool_items]
    assert ops_tool_slugs[:4] == [
        "adb-command-search",
        "linux-command-search",
        "git-command-search",
        "docker-command-search",
    ]
    assert "text-format-cleaner" in ops_tool_slugs
    assert "invisible-control-chars" in ops_tool_slugs
    command_tool_items = [item for item in ops_tool_items if item["component"] == "ToolCommandCatalog"]
    assert all(item["input_mode"] == "local" for item in command_tool_items)
    assert any(
        item["slug"] == "local-ip-lookup"
        and item["input_mode"] == "form"
        and item["component"] == "ToolLocalIpLookup"
        for item in ops_tool_items
    )

    local_tools = client.get("/api/tools", params={"category": "game"})
    assert local_tools.status_code == 200
    assert any(item["slug"] == "dino-runner" for item in local_tools.json()["data"])

    other_tools = client.get("/api/tools", params={"category": "other"})
    assert other_tools.status_code == 200
    other_tool_items = other_tools.json()["data"]
    other_tool_slugs = [item["slug"] for item in other_tool_items]
    assert other_tool_slugs[:4] == [
        "market-quote",
        "today-stock-index",
        "today-gold-price",
        "today-oil-price",
    ]
    assert "market-quote" in other_tool_slugs
    assert "today-oil-price" in other_tool_slugs
    assert "today-international-crude" in other_tool_slugs
    assert "today-gold-price" in other_tool_slugs
    assert "today-exchange-rate" in other_tool_slugs
    assert "today-food-price" in other_tool_slugs
    assert "today-silver-price" in other_tool_slugs
    assert "today-stock-index" in other_tool_slugs
    assert "today-building-materials" in other_tool_slugs
    assert "douyin-id-extractor" in other_tool_slugs
    assert "abstract-fan" in other_tool_slugs
    assert "abstract-ac" in other_tool_slugs
    assert any(item["slug"] == "market-quote" and item["component"] == "ToolMarketQuote" for item in other_tool_items)
    assert any(item["slug"] == "abstract-fan" and item["component"] == "ToolAbstractAppliance" for item in other_tool_items)

    image_tools = client.get("/api/tools", params={"category": "image"})
    assert image_tools.status_code == 200
    image_tool_items = image_tools.json()["data"]
    image_tool_slugs = [item["slug"] for item in image_tool_items]
    compressor_index = image_tool_slugs.index("image-compressor")
    assert image_tool_slugs[compressor_index + 1] == "image-upscaler"
    compressor_params = image_tool_items[compressor_index]["params"]
    assert compressor_params[0]["key"] == "mode"
    assert [item["value"] for item in compressor_params[0]["options"]] == ["byte_size", "dimensions"]
    assert compressor_params[1]["key"] == "target_kb"
    assert compressor_params[2]["key"] == "scale"
    assert image_tool_items[compressor_index + 1]["name"] == "图片增大"
    upscaler_params = image_tool_items[compressor_index + 1]["params"]
    assert upscaler_params[0]["key"] == "mode"
    assert [item["value"] for item in upscaler_params[0]["options"]] == ["byte_size", "dimensions"]


def test_price_snapshot_tools_with_mocked_public_sources(monkeypatch) -> None:
    from app.tools import market_quote, price_snapshot

    def fake_json(url, params=None, headers=None):
        if url == price_snapshot.XXAPI_OIL_URL:
            return {
                "code": 200,
                "data": [
                    {
                        "regionName": "北京市",
                        "n89": "7.18",
                        "n92": "7.66",
                        "n95": "8.16",
                        "n98": "9.14",
                        "n0": "7.37",
                        "date": "2026-06-05",
                    }
                ],
            }
        if url == price_snapshot.FRANKFURTER_URL:
            if params and params.get("from") == "USD" and params.get("to") == "CNY":
                return {"date": "2026-06-04", "rates": {"CNY": 7.1}}
            return {
                "date": "2026-06-04",
                "rates": {"CNY": 7.1, "EUR": 0.86, "JPY": 143.1, "GBP": 0.74, "HKD": 7.83},
            }
        raise AssertionError(f"unexpected json url: {url}")

    def fake_text(url, params=None, headers=None, encoding="utf-8"):
        if url.startswith(price_snapshot.SINA_HQ_URL):
            return "\n".join(
                [
                    'var hq_str_hf_CL="93.149,,93.050,93.080,93.540,92.520,12:29:36,93.040,92.820,0,3,6,2026-06-05,纽约原油,0";',
                    'var hq_str_hf_OIL="95.418,,95.310,95.340,95.900,94.790,12:29:36,95.030,95.290,0,3,2,2026-06-05,布伦特原油,9387";',
                    'var hq_str_hf_GC="4465.205,,4465.900,4466.300,4508.700,4461.200,12:29:33,4505.000,4503.000,0,1,2,2026-06-05,纽约黄金,0";',
                    'var hq_str_hf_SI="72.796,,72.790,72.820,74.380,72.520,12:29:21,73.971,74.185,0,2,1,2026-06-05,纽约白银,0";',
                    'var hq_str_USDCNY="12:23:09,7.1000,7.1000,7.1000,86,7.1000,7.1000,7.1000,7.1000,美元人民币,2026-06-05";',
                ]
            )
        if url == price_snapshot.MOA_LIST_URL:
            return '<a href="./202606/t20260604_6484728.htm">6月4日：“农产品批发价格200指数”比昨天上升0.12个点2026-06-04</a>'
        if "t20260604_6484728.htm" in url:
            return (
                "猪肉平均价格为14.73元/公斤，比昨天下降0.1%；"
                "鸡蛋10.48元/公斤，比昨天上升0.4%；"
                "白条鸡17.29元/公斤，比昨天上升0.9%；"
                "28种蔬菜平均价格为4.26元/公斤，比昨天上升1.2%。"
            )
        if url == price_snapshot.MYSTEEL_MOBILE_URL:
            return '<a href="/x">6月5日(12:10)南京市场建筑钢材价格行情 高线 螺纹钢 盘螺 06-05</a>'
        if url in price_snapshot.MYSTEEL_SAND_STONE_URLS:
            return '<li>6月5日南京市场建设用砂石价格行情 河砂 机制砂 碎石 2026-06-05 11:30</li>'
        raise AssertionError(f"unexpected text url: {url}")

    monkeypatch.setattr(price_snapshot, "_request_json", fake_json)
    monkeypatch.setattr(price_snapshot, "_request_text", fake_text)
    monkeypatch.setattr(
        market_quote,
        "run",
        lambda query, market="auto": json.dumps(
            {
                "cards": [
                    {
                        "name": "上证指数",
                        "symbol": "000001",
                        "latest": 3384.5,
                        "change_percent": 0.35,
                        "updated_at": "2026-06-05 12:30:00",
                    }
                ]
            },
            ensure_ascii=False,
        ),
    )

    assert "今日油价 - 北京市" in price_snapshot.run_domestic_oil("北京")
    assert "WTI/纽约原油" in price_snapshot.run_international_crude()
    assert "纽约黄金折算" in price_snapshot.run_gold()
    assert "USD/CNY" in price_snapshot.run_exchange_rate("USD")
    assert "100 USD -> CNY" in price_snapshot.run_exchange_rate("100美元兑人民币")
    assert "猪肉" in price_snapshot.run_food_price()
    assert "纽约白银折算" in price_snapshot.run_silver()
    assert "上证指数(000001)" in price_snapshot.run_stock_index("")
    building_materials = price_snapshot.run_building_materials()
    assert "南京市场建筑钢材价格行情" in building_materials
    assert "南京市场建设用砂石价格行情" in building_materials
    assert "红砖：0.6-0.8 元/块" in building_materials


def test_local_ip_lookup_helpers() -> None:
    from app.tools import local_ip_lookup

    assert local_ip_lookup.extract_ip_from_text("当前 IP：1.2.3.4 来自于：中国") == "1.2.3.4"
    assert local_ip_lookup.extract_ip_from_text("bad 999.999.999.999 2001:db8::1") == "2001:db8::1"
    same = local_ip_lookup.analyze_split_ip("1.2.3.4", "1.2.3.4")
    split = local_ip_lookup.analyze_split_ip("1.2.3.4", "5.6.7.8")
    unknown = local_ip_lookup.analyze_split_ip("", "5.6.7.8")
    assert same["status"] == "same"
    assert split["status"] == "split"
    assert unknown["status"] == "unknown"


def test_market_quote_helpers() -> None:
    import httpx

    from app.tools import market_quote

    assert ("1.600519", "沪市") in market_quote._eastmoney_candidates("600519", "auto")
    assert ("116.00700", "港股") in market_quote._eastmoney_candidates("HK00700", "auto")
    assert ("105.AAPL", "美股") in market_quote._eastmoney_candidates("AAPL", "auto")
    assert market_quote._sina_cn_symbol("2.931787") == "si931787"

    parsed = market_quote._parse_fund_response(
        'jsonpgz({"fundcode":"110022","name":"易方达消费行业股票","jzrq":"2026-06-02","dwjz":"2.8930","gsz":"2.8491","gszzl":"-1.52","gztime":"2026-06-03 10:56"});'
    )
    assert parsed["fundcode"] == "110022"
    assert parsed["gszzl"] == "-1.52"

    robot_text = market_quote._build_robot_text(
        [
            {
                "name": "贵州茅台",
                "symbol": "600519",
                "latest": 1281.01,
                "currency": "CNY",
                "change_percent": -2.01,
                "change": -26.21,
                "updated_at": "2026-06-03 10:59:32",
                "source": "东方财富公开行情",
            }
        ],
        [],
    )
    assert "贵州茅台(600519)" in robot_text
    assert "东方财富公开行情" in robot_text

    class FakeResponse:
        def __init__(self, status_code: int, payload: object, text: str | None = None) -> None:
            self.status_code = status_code
            self._payload = payload
            self._text = text if text is not None else json.dumps(payload, ensure_ascii=False)
            self.content = self._text.encode("gbk")

        def raise_for_status(self) -> None:
            if self.status_code < 400:
                return
            request = httpx.Request("GET", "https://example.test")
            response = httpx.Response(self.status_code, request=request)
            raise httpx.HTTPStatusError("mock error", request=request, response=response)

        @property
        def text(self) -> str:
            return self._text

        def json(self) -> object:
            return self._payload

    class FakeClient:
        def __init__(self) -> None:
            self.urls: list[str] = []

        def get(self, url: str, **_: object) -> FakeResponse:
            self.urls.append(url)
            if url == market_quote.EASTMONEY_URL:
                return FakeResponse(502, {})
            if url == market_quote.EASTMONEY_TRENDS_URL:
                return FakeResponse(
                    200,
                    {
                        "data": {
                            "preClose": 7019.03,
                            "time": 1780474296,
                            "trends": [
                                "2026-06-03 09:30,6987.87,6987.87,6987.87,6987.87,100,1000.00,6987.870",
                                "2026-06-03 09:31,6987.87,6937.71,6990.00,6937.71,200,2000.00,6960.120",
                            ],
                        }
                    },
                )
            return FakeResponse(
                200,
                {
                    "data": {
                        "diff": [
                            {
                                "f1": 2,
                                "f2": 693771,
                                "f3": -116,
                                "f4": -8132,
                                "f5": 2348447,
                                "f6": 15515966819.73,
                                "f7": 218,
                                "f8": 138,
                                "f12": "399997",
                                "f13": 0,
                                "f14": "中证白酒",
                                "f15": 699887,
                                "f16": 684578,
                                "f17": 698787,
                                "f18": 701903,
                                "f20": 2472012797975,
                                "f21": 2471856071232,
                                "f124": 1780474296,
                                "f152": 2,
                            }
                        ]
                    }
                },
            )

    fake_client = FakeClient()
    fallback_card, fallback_error = market_quote._query_eastmoney("399997", "auto", fake_client)  # type: ignore[arg-type]
    assert fallback_error is None
    assert fallback_card["name"] == "中证白酒"
    assert fallback_card["latest"] == 6937.71
    assert fallback_card["chart"]["time_range"] == "2026-06-03 09:30 - 2026-06-03 09:31"
    assert fallback_card["chart"]["points"] == [
        ["2026-06-03 09:30", 6987.87, 6987.87, 100.0],
        ["2026-06-03 09:31", 6937.71, 6960.12, 200.0],
    ]
    assert fake_client.urls == [market_quote.EASTMONEY_URL, market_quote.EASTMONEY_ULIST_URL, market_quote.EASTMONEY_TRENDS_URL]

    class IndexFallbackClient:
        def __init__(self) -> None:
            self.requests: list[tuple[str, str]] = []

        def get(self, url: str, **kwargs: object) -> FakeResponse:
            params = kwargs.get("params")
            params = params if isinstance(params, dict) else {}
            secid = str(params.get("secid") or params.get("secids") or "")
            self.requests.append((url, secid))
            if url == market_quote.EASTMONEY_URL and secid == "1.931787":
                return FakeResponse(200, {"data": None})
            if url == market_quote.EASTMONEY_ULIST_URL and secid == "1.931787":
                return FakeResponse(200, {"data": None})
            if url == market_quote.EASTMONEY_URL and secid == "2.931787":
                return FakeResponse(502, {})
            if url == market_quote.EASTMONEY_TRENDS_URL and secid == "2.931787":
                return FakeResponse(
                    200,
                    {
                        "data": {
                            "preClose": 1140.61,
                            "time": 1780537535,
                            "trends": [
                                "2026-06-04 09:30,1144.48,1144.48,1144.48,1143.04,32088,42811410.00,1143.023",
                                "2026-06-04 09:31,1144.31,1153.13,1153.36,1144.31,91280,156748134.00,1145.059",
                            ],
                        }
                    },
                )
            return FakeResponse(
                200,
                {
                    "data": {
                        "diff": [
                            {
                                "f1": 2,
                                "f2": 113631,
                                "f3": -38,
                                "f4": -430,
                                "f5": 964138,
                                "f6": 1370808067.7,
                                "f7": 188,
                                "f8": 11,
                                "f12": "931787",
                                "f13": 2,
                                "f14": "港股创新药",
                                "f15": 115540,
                                "f16": 113395,
                                "f17": 114448,
                                "f18": 114061,
                                "f20": 1302757542438,
                                "f21": 1302757542438,
                                "f124": 1780537535,
                                "f152": 2,
                            }
                        ]
                    }
                },
            )

    index_client = IndexFallbackClient()
    index_card, index_error = market_quote._query_eastmoney("931787", "auto", index_client)  # type: ignore[arg-type]
    assert index_error is None
    assert index_card["name"] == "港股创新药"
    assert index_card["asset_type"] == "指数"
    assert index_card["latest"] == 1136.31
    assert index_card["chart"]["pre_close"] == 1140.61
    assert index_card["chart"]["trade_date"] == "2026-06-04"
    assert index_card["chart"]["points"][1] == ["2026-06-04 09:31", 1153.13, 1145.059, 91280.0]
    assert index_client.requests == [
        (market_quote.EASTMONEY_URL, "1.931787"),
        (market_quote.EASTMONEY_ULIST_URL, "1.931787"),
        (market_quote.EASTMONEY_TRENDS_URL, "1.931787"),
        (market_quote.EASTMONEY_KLINE_URL, "1.931787"),
        (market_quote.SINA_KLINE_URL, ""),
        (market_quote.EASTMONEY_URL, "2.931787"),
        (market_quote.EASTMONEY_ULIST_URL, "2.931787"),
        (market_quote.EASTMONEY_TRENDS_URL, "2.931787"),
    ]

    class TrendsOnlyClient:
        def __init__(self) -> None:
            self.requests: list[tuple[str, str]] = []

        def get(self, url: str, **kwargs: object) -> FakeResponse:
            params = kwargs.get("params")
            params = params if isinstance(params, dict) else {}
            secid = str(params.get("secid") or params.get("secids") or "")
            self.requests.append((url, secid))
            if url in {market_quote.EASTMONEY_URL, market_quote.EASTMONEY_ULIST_URL}:
                return FakeResponse(502, {})
            return FakeResponse(
                200,
                {
                    "data": {
                        "code": "399997",
                        "name": "中证白酒",
                        "market": 0,
                        "preClose": 6937.71,
                        "time": 1780537535,
                        "trends": [
                            "2026-06-04 09:30,6889.51,6889.51,6889.51,6889.51,100,1000.00,6889.510",
                            "2026-06-04 09:31,6889.51,6899.51,6899.51,6889.51,120,1200.00,6894.510",
                        ],
                    }
                },
            )

    trends_client = TrendsOnlyClient()
    trends_card, trends_error = market_quote._query_eastmoney("399997", "auto", trends_client)  # type: ignore[arg-type]
    assert trends_error is None
    assert trends_card["name"] == "中证白酒"
    assert trends_card["latest"] == 6899.51
    assert trends_card["change"] == -38.2
    assert trends_card["chart"]["kind"] == "intraday"
    assert trends_client.requests == [
        (market_quote.EASTMONEY_URL, "0.399997"),
        (market_quote.EASTMONEY_ULIST_URL, "0.399997"),
        (market_quote.EASTMONEY_TRENDS_URL, "0.399997"),
    ]

    class KlineOnlyClient:
        def __init__(self) -> None:
            self.requests: list[tuple[str, str]] = []

        def get(self, url: str, **kwargs: object) -> FakeResponse:
            params = kwargs.get("params")
            params = params if isinstance(params, dict) else {}
            secid = str(params.get("secid") or params.get("secids") or "")
            self.requests.append((url, secid))
            if url in {market_quote.EASTMONEY_URL, market_quote.EASTMONEY_ULIST_URL}:
                return FakeResponse(502, {})
            if url == market_quote.EASTMONEY_TRENDS_URL:
                return FakeResponse(200, {"data": {}})
            return FakeResponse(
                200,
                {
                    "data": {
                        "code": "399997",
                        "name": "中证白酒",
                        "market": 0,
                        "klines": [
                            "2026-06-01,7000.00,7010.00,7030.00,6990.00,1000,1000000.00,1.00,0.10,7.00,0.10",
                            "2026-06-02,7010.00,6990.00,7020.00,6980.00,1200,1200000.00,1.00,-0.29,-20.00,0.12",
                            "2026-06-03,6990.00,6899.51,7000.00,6880.00,1400,1400000.00,1.00,-1.29,-90.49,0.14",
                        ],
                    }
                },
            )

    kline_client = KlineOnlyClient()
    kline_card, kline_error = market_quote._query_eastmoney("399997", "auto", kline_client)  # type: ignore[arg-type]
    assert kline_error is None
    assert kline_card["name"] == "中证白酒"
    assert kline_card["latest"] == 6899.51
    assert kline_card["previous_close"] == 6990.0
    assert kline_card["chart"]["kind"] == "daily_kline"
    assert kline_card["chart"]["points"][-1] == ["2026-06-03", 6899.51, 6966.5033, 1400.0]
    assert kline_client.requests == [
        (market_quote.EASTMONEY_URL, "0.399997"),
        (market_quote.EASTMONEY_ULIST_URL, "0.399997"),
        (market_quote.EASTMONEY_TRENDS_URL, "0.399997"),
        (market_quote.EASTMONEY_KLINE_URL, "0.399997"),
    ]

    class SinaOnlyClient:
        def __init__(self) -> None:
            self.requests: list[tuple[str, str]] = []

        def get(self, url: str, **kwargs: object) -> FakeResponse:
            params = kwargs.get("params")
            params = params if isinstance(params, dict) else {}
            query_id = str(params.get("secid") or params.get("secids") or params.get("symbol") or "")
            self.requests.append((url, query_id))
            if url in {
                market_quote.EASTMONEY_URL,
                market_quote.EASTMONEY_ULIST_URL,
                market_quote.EASTMONEY_TRENDS_URL,
                market_quote.EASTMONEY_KLINE_URL,
            }:
                return FakeResponse(200, {"data": {}})
            if url == market_quote.SINA_KLINE_URL:
                return FakeResponse(
                    200,
                    [
                        {
                            "day": "2026-06-04 09:30:00",
                            "open": "6889.511",
                            "high": "6891.000",
                            "low": "6880.000",
                            "close": "6889.511",
                            "volume": "1000",
                            "amount": "1000000.00",
                        },
                        {
                            "day": "2026-06-04 09:35:00",
                            "open": "6889.511",
                            "high": "6901.000",
                            "low": "6881.000",
                            "close": "6899.511",
                            "volume": "1200",
                            "amount": "1200000.00",
                        },
                    ],
                )
            return FakeResponse(
                200,
                {},
                'var hq_str_sz399997="中证白酒,6889.511,6937.715,6899.511,6901.000,6880.000,0.000,0.000,2200,2200000.00,0,0.000,0,0.000,0,0.000,0,0.000,0,0.000,0,0.000,0,0.000,0,0.000,0,0.000,0,0.000,2026-06-04,09:35:00,00";',
            )

    sina_client = SinaOnlyClient()
    sina_card, sina_error = market_quote._query_eastmoney("399997", "auto", sina_client)  # type: ignore[arg-type]
    assert sina_error is None
    assert sina_card["name"] == "中证白酒"
    assert sina_card["latest"] == 6899.511
    assert sina_card["previous_close"] == 6937.715
    assert sina_card["provider"] == "sina"
    assert sina_card["chart"]["kind"] == "intraday_kline"
    assert sina_card["chart"]["points"][-1] == ["2026-06-04 09:35:00", 6899.511, 6894.511, 1200.0]
    assert sina_client.requests == [
        (market_quote.EASTMONEY_URL, "0.399997"),
        (market_quote.EASTMONEY_ULIST_URL, "0.399997"),
        (market_quote.EASTMONEY_TRENDS_URL, "0.399997"),
        (market_quote.EASTMONEY_KLINE_URL, "0.399997"),
        (market_quote.SINA_KLINE_URL, "sz399997"),
        (f"{market_quote.SINA_HQ_URL}sz399997", ""),
    ]

    class MixedDateSinaClient(SinaOnlyClient):
        def get(self, url: str, **kwargs: object) -> FakeResponse:
            response = super().get(url, **kwargs)
            if url != market_quote.SINA_KLINE_URL:
                return response
            return FakeResponse(
                200,
                [
                    {
                        "day": "2026-06-03 09:40:00",
                        "open": "6870.000",
                        "high": "6880.000",
                        "low": "6860.000",
                        "close": "6872.146",
                        "volume": "1000",
                    },
                    {
                        "day": "2026-06-03 09:45:00",
                        "open": "6872.146",
                        "high": "6885.000",
                        "low": "6870.000",
                        "close": "6882.404",
                        "volume": "1200",
                    },
                    {
                        "day": "2026-06-04 14:55:00",
                        "open": "6760.000",
                        "high": "6770.000",
                        "low": "6750.000",
                        "close": "6762.372",
                        "volume": "1400",
                    },
                    {
                        "day": "2026-06-04 15:00:00",
                        "open": "6762.372",
                        "high": "6768.000",
                        "low": "6758.000",
                        "close": "6758.918",
                        "volume": "1600",
                    },
                    {
                        "day": "2026-06-05 09:35:00",
                        "open": "6860.000",
                        "high": "6870.000",
                        "low": "6850.000",
                        "close": "6866.545",
                        "volume": "2400",
                    },
                ],
            )

    mixed_card, mixed_error = market_quote._query_eastmoney("399997", "auto", MixedDateSinaClient())  # type: ignore[arg-type]
    assert mixed_error is None
    assert mixed_card["chart"]["trade_date"] == "2026-06-04"
    assert mixed_card["chart"]["display_note"] == "当前交易日K线数据不足，展示上一交易日"
    assert {point[0][:10] for point in mixed_card["chart"]["points"]} == {"2026-06-04"}
    assert mixed_card["chart"]["time_range"] == "2026-06-04 14:55:00 - 2026-06-04 15:00:00"


def test_douyin_id_extractor_execute_offline() -> None:
    sec_uid = "MS4wLjABAAAAabcdef1234567890"
    response = client.post(
        "/api/tools/douyin-id-extractor/execute",
        json={
            "text": f"https://www.douyin.com/user/{sec_uid}?uid=123456789&unique_id=qixiang_tools",
            "params": {"resolve_links": False},
        },
    )
    assert response.status_code == 200
    result = response.json()["data"]["result"]
    assert sec_uid in result
    assert "123456789" in result
    assert "qixiang_tools" in result


def test_json_format_execute() -> None:
    response = client.post(
        "/api/tools/json-format/execute",
        json={"text": '{"b":1,"a":2}', "params": {"indent": 2}},
    )
    assert response.status_code == 200
    assert '"b": 1' in response.json()["data"]["result"]


def test_json_format_accepts_json_like_text() -> None:
    response = client.post(
        "/api/tools/json-format/execute",
        json={
            "text": """
[Bridge]
但人总是这样
明知道水开了
{
  title: 'City Lights',
  tags: '华语流行情歌',
  mv: 'chirp-crow',
}
""",
            "params": {"indent": 2},
        },
    )
    assert response.status_code == 200
    result = response.json()["data"]["result"]
    assert '"title": "City Lights"' in result
    assert '"mv": "chirp-crow"' in result


def test_json_format_accepts_object_without_braces() -> None:
    response = client.post(
        "/api/tools/json-format/execute",
        json={
            "text": """
title: 'City Lights',
artist: '琦湘',
count: 2
""",
            "params": {"indent": 2},
        },
    )
    assert response.status_code == 200
    result = response.json()["data"]["result"]
    assert '"title": "City Lights"' in result
    assert '"artist": "琦湘"' in result
    assert '"count": 2' in result


def test_json_format_repairs_multiline_string_field() -> None:
    response = client.post(
        "/api/tools/json-format/execute",
        json={
            "text": """
{
  "lyrics": "[Chorus]
可惜没如果
只有那分跳闸的结果
热腾腾的面
变成一锅冰凉的辜负",
  "title": "可惜没如果",
  "artist": "林俊杰"
}
""",
            "params": {"action": "format", "indent": 2},
        },
    )
    assert response.status_code == 200
    result = response.json()["data"]["result"]
    assert '"title": "可惜没如果"' in result
    assert '"artist": "林俊杰"' in result
    assert "\\n可惜没如果\\n只有那分跳闸的结果" in result


def test_json_format_strict_validate_success() -> None:
    response = client.post(
        "/api/tools/json-format/execute",
        json={
            "text": '{"title":"City Lights","count":2}',
            "params": {"action": "validate"},
        },
    )
    assert response.status_code == 200
    result = response.json()["data"]["result"]
    assert "校验结果：通过" in result
    assert '"title": "City Lights"' in result


def test_json_format_strict_validate_failure() -> None:
    response = client.post(
        "/api/tools/json-format/execute",
        json={
            "text": "title: 'City Lights'",
            "params": {"action": "validate"},
        },
    )
    assert response.status_code == 200
    result = response.json()["data"]["result"]
    assert "校验结果：不通过" in result
    assert "原因：" in result


def test_base64_tool_execute() -> None:
    response = client.post(
        "/api/tools/base64-codec/execute",
        json={"text": "hello", "params": {"action": "encode"}},
    )
    assert response.status_code == 200
    assert response.json()["data"]["result"] == "aGVsbG8="


def test_image_to_base64_upload() -> None:
    png_bytes = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO0pS1kAAAAASUVORK5CYII="
    )
    response = client.post(
        "/api/tools/image-to-base64/upload",
        data={"params": "{}"},
        files={"file": ("tiny.png", png_bytes, "image/png")},
    )
    assert response.status_code == 200
    assert response.json()["data"]["result"].startswith("data:image/png;base64,")


def test_image_tools_upload_and_execute() -> None:
    png_bytes = make_png_bytes()
    noise_png_bytes = make_noise_png_bytes()

    compress_response = client.post(
        "/api/tools/image-compressor/upload",
        data={"params": json.dumps({"quality": 60, "output_format": "jpg", "max_width": 80, "max_height": 80})},
        files={"file": ("sample.png", png_bytes, "image/png")},
    )
    assert compress_response.status_code == 200
    assert "attachment" in compress_response.headers.get("content-disposition", "")
    assert compress_response.headers["content-type"].startswith("image/jpeg")
    assert compress_response.headers["access-control-expose-headers"] == "Content-Disposition"

    target_kb = 20
    byte_compress_response = client.post(
        "/api/tools/image-compressor/upload",
        data={"params": json.dumps({"mode": "byte_size", "target_kb": target_kb})},
        files={"file": ("noise.png", noise_png_bytes, "image/png")},
    )
    assert byte_compress_response.status_code == 200
    assert byte_compress_response.headers["content-type"].startswith("image/jpeg")
    assert len(byte_compress_response.content) <= target_kb * 1024

    dimension_compress_response = client.post(
        "/api/tools/image-compressor/upload",
        data={"params": json.dumps({"mode": "dimensions", "scale": 4})},
        files={"file": ("noise.png", noise_png_bytes, "image/png")},
    )
    assert dimension_compress_response.status_code == 200
    assert dimension_compress_response.headers["content-type"].startswith("image/png")
    compressed_image = Image.open(io.BytesIO(dimension_compress_response.content))
    assert compressed_image.size == (120, 120)

    analyze_response = client.post(
        "/api/tools/image-color-analyzer/upload",
        data={"params": json.dumps({"top_n": 3})},
        files={"file": ("sample.png", png_bytes, "image/png")},
    )
    assert analyze_response.status_code == 200
    assert "#FF0000" in analyze_response.json()["data"]["result"]

    watermarked_response = client.post(
        "/api/tools/image-watermarker/upload",
        data={"params": json.dumps({"text": "天问", "position": "bottom_right", "opacity": 45, "font_size": 24, "color": "#ffffff"})},
        files={"file": ("sample.png", png_bytes, "image/png")},
    )
    assert watermarked_response.status_code == 200
    assert "attachment" in watermarked_response.headers.get("content-disposition", "")

    grid_response = client.post(
        "/api/tools/image-grid-splitter/upload",
        data={"params": json.dumps({"rows": 3, "cols": 3})},
        files={"file": ("sample.png", png_bytes, "image/png")},
    )
    assert grid_response.status_code == 200
    assert "attachment" in grid_response.headers.get("content-disposition", "")


def test_file_result_execute_exposes_filename_header() -> None:
    response = client.post(
        "/api/tools/handwritten-signature/execute",
        json={
            "text": "琦湘",
            "params": {"color": "#111111", "font_size": 120, "transparent_bg": True},
        },
    )
    assert response.status_code == 200
    assert "attachment" in response.headers.get("content-disposition", "")
    assert response.headers["content-type"].startswith("image/png")
    assert response.headers["access-control-expose-headers"] == "Content-Disposition"


def test_image_compressor_accepts_heic_upload() -> None:
    response = client.post(
        "/api/tools/image-compressor/upload",
        data={"params": json.dumps({"mode": "dimensions", "scale": 2})},
        files={"file": ("sample.heic", make_heic_bytes(), "application/octet-stream")},
    )
    assert response.status_code == 200
    assert "attachment" in response.headers.get("content-disposition", "")
    assert response.headers["content-type"].startswith("image/heic")


def test_qr_barcode_and_gif_tools_execute() -> None:
    qr_generate_response = client.post(
        "/api/tools/qr-code-generator/execute",
        json={
            "text": "https://example.com",
            "params": {"box_size": 6, "border": 2, "fill_color": "#000000", "back_color": "#ffffff"},
        },
    )
    assert qr_generate_response.status_code == 200
    assert qr_generate_response.headers["content-type"].startswith("image/png")

    qr_stream = io.BytesIO()
    qrcode.make("https://example.com").save(qr_stream, format="PNG")
    qr_response = client.post(
        "/api/tools/qr-code-decoder/upload",
        data={"params": "{}"},
        files={"file": ("qr.png", qr_stream.getvalue(), "image/png")},
    )
    assert qr_response.status_code == 200
    assert qr_response.json()["data"]["result"] == "https://example.com"

    barcode_response = client.post(
        "/api/tools/barcode-generator/execute",
        json={"text": "TW-20260403", "params": {"barcode_type": "code128"}},
    )
    assert barcode_response.status_code == 200
    assert barcode_response.headers["content-type"].startswith("image/png")

    gif_bytes = make_gif_bytes()
    gif_split_response = client.post(
        "/api/tools/gif-splitter/upload",
        data={"params": "{}"},
        files={"file": ("demo.gif", gif_bytes, "image/gif")},
    )
    assert gif_split_response.status_code == 200
    assert "attachment" in gif_split_response.headers.get("content-disposition", "")

    gif_scale_response = client.post(
        "/api/tools/gif-scaler/upload",
        data={"params": json.dumps({"scale_percent": 50})},
        files={"file": ("demo.gif", gif_bytes, "image/gif")},
    )
    assert gif_scale_response.status_code == 200
    assert gif_scale_response.headers["content-type"].startswith("image/gif")


def test_image_zip_and_office_extract_tools() -> None:
    png_bytes = make_png_bytes((0, 255, 0), (80, 80))
    blue_png = make_png_bytes((0, 0, 255), (80, 80))
    image_zip = make_zip_bytes(
        {
            "a.png": png_bytes,
            "b.png": blue_png,
        }
    )

    merge_response = client.post(
        "/api/tools/image-merger/upload",
        data={"params": json.dumps({"direction": "horizontal", "gap": 8, "output_format": "png"})},
        files={"file": ("images.zip", image_zip, "application/zip")},
    )
    assert merge_response.status_code == 200
    assert merge_response.headers["content-type"].startswith("image/png")

    gif_make_response = client.post(
        "/api/tools/gif-maker/upload",
        data={"params": json.dumps({"duration": 120, "loop": 0})},
        files={"file": ("images.zip", image_zip, "application/zip")},
    )
    assert gif_make_response.status_code == 200
    assert gif_make_response.headers["content-type"].startswith("image/gif")

    gif_zip = make_zip_bytes({"a.gif": make_gif_bytes(), "b.gif": make_gif_bytes()})
    gif_merge_response = client.post(
        "/api/tools/gif-merger/upload",
        data={"params": "{}"},
        files={"file": ("gifs.zip", gif_zip, "application/zip")},
    )
    assert gif_merge_response.status_code == 200
    assert gif_merge_response.headers["content-type"].startswith("image/gif")

    docx_bytes = make_zip_bytes({"word/media/image1.png": png_bytes})
    word_response = client.post(
        "/api/tools/word-image-extractor/upload",
        data={"params": "{}"},
        files={"file": ("demo.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert word_response.status_code == 200
    assert "attachment" in word_response.headers.get("content-disposition", "")

    xlsx_bytes = make_zip_bytes({"xl/media/image1.png": png_bytes})
    excel_response = client.post(
        "/api/tools/excel-image-extractor/upload",
        data={"params": "{}"},
        files={"file": ("demo.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert excel_response.status_code == 200
    assert "attachment" in excel_response.headers.get("content-disposition", "")

    pptx_bytes = make_zip_bytes({"ppt/media/image1.png": png_bytes})
    ppt_response = client.post(
        "/api/tools/ppt-image-extractor/upload",
        data={"params": "{}"},
        files={"file": ("demo.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert ppt_response.status_code == 200
    assert "attachment" in ppt_response.headers.get("content-disposition", "")


def test_document_conversion_tools_upload() -> None:
    docx_response = client.post(
        "/api/tools/docx-to-pdf/upload",
        data={"params": "{}"},
        files={
            "file": (
                "demo.docx",
                make_docx_bytes("Hello from Word"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert docx_response.status_code == 200
    assert docx_response.headers["content-type"].startswith("application/pdf")

    pdf_bytes = make_pdf_bytes("Hello from PDF")
    pdf_to_word_response = client.post(
        "/api/tools/pdf-to-word/upload",
        data={"params": "{}"},
        files={"file": ("demo.pdf", pdf_bytes, "application/pdf")},
    )
    assert pdf_to_word_response.status_code == 200
    assert "wordprocessingml.document" in pdf_to_word_response.headers["content-type"]

    pdf_to_jpg_response = client.post(
        "/api/tools/pdf-to-jpg/upload",
        data={"params": json.dumps({"scale": 2, "quality": 90})},
        files={"file": ("demo.pdf", pdf_bytes, "application/pdf")},
    )
    assert pdf_to_jpg_response.status_code == 200
    assert pdf_to_jpg_response.headers["content-type"].startswith("application/zip")
    with zipfile.ZipFile(io.BytesIO(pdf_to_jpg_response.content)) as archive:
        assert any(name.endswith(".jpg") for name in archive.namelist())

    pdf_to_epub_response = client.post(
        "/api/tools/pdf-to-epub/upload",
        data={"params": json.dumps({"title": "PDF 电子书"})},
        files={"file": ("demo.pdf", pdf_bytes, "application/pdf")},
    )
    assert pdf_to_epub_response.status_code == 200
    assert "application/epub+zip" in pdf_to_epub_response.headers["content-type"]


def test_epub_and_image_to_pdf_upload() -> None:
    epub_response = client.post(
        "/api/tools/epub-to-pdf/upload",
        data={"params": json.dumps({"title": "EPUB 导出 PDF"})},
        files={"file": ("demo.epub", make_epub_bytes("电子书标题", "Hello from EPUB"), "application/epub+zip")},
    )
    assert epub_response.status_code == 200
    assert epub_response.headers["content-type"].startswith("application/pdf")

    jpg_response = client.post(
        "/api/tools/jpg-to-pdf/upload",
        data={"params": "{}"},
        files={"file": ("demo.jpg", make_jpg_bytes((80, 120, 200), (180, 120)), "image/jpeg")},
    )
    assert jpg_response.status_code == 200
    assert jpg_response.headers["content-type"].startswith("application/pdf")


def test_more_document_conversion_tools_upload() -> None:
    pdf_bytes = make_pdf_bytes("Hello from PDF")
    docx_bytes = make_docx_bytes("Hello from Word")
    txt_bytes = make_txt_bytes("第一行\n第二行\n第三行")

    pdf_to_txt_response = client.post(
        "/api/tools/pdf-to-txt/upload",
        data={"params": "{}"},
        files={"file": ("demo.pdf", pdf_bytes, "application/pdf")},
    )
    assert pdf_to_txt_response.status_code == 200
    assert pdf_to_txt_response.headers["content-type"].startswith("text/plain")

    word_to_txt_response = client.post(
        "/api/tools/word-to-txt/upload",
        data={"params": "{}"},
        files={
            "file": (
                "demo.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert word_to_txt_response.status_code == 200
    assert word_to_txt_response.headers["content-type"].startswith("text/plain")

    txt_to_word_response = client.post(
        "/api/tools/txt-to-word/upload",
        data={"params": json.dumps({"title": "TXT 文档"})},
        files={"file": ("demo.txt", txt_bytes, "text/plain")},
    )
    assert txt_to_word_response.status_code == 200
    assert "wordprocessingml.document" in txt_to_word_response.headers["content-type"]

    txt_to_pdf_response = client.post(
        "/api/tools/txt-to-pdf/upload",
        data={"params": json.dumps({"title": "TXT PDF"})},
        files={"file": ("demo.txt", txt_bytes, "text/plain")},
    )
    assert txt_to_pdf_response.status_code == 200
    assert txt_to_pdf_response.headers["content-type"].startswith("application/pdf")

    txt_to_epub_response = client.post(
        "/api/tools/txt-to-epub/upload",
        data={"params": json.dumps({"title": "TXT EPUB"})},
        files={"file": ("demo.txt", txt_bytes, "text/plain")},
    )
    assert txt_to_epub_response.status_code == 200
    assert "application/epub+zip" in txt_to_epub_response.headers["content-type"]

    heic_to_pdf_response = client.post(
        "/api/tools/heic-to-pdf/upload",
        data={"params": "{}"},
        files={"file": ("demo.heic", make_heic_bytes(), "image/heic")},
    )
    assert heic_to_pdf_response.status_code == 200
    assert heic_to_pdf_response.headers["content-type"].startswith("application/pdf")


def test_even_more_document_conversion_tools_upload() -> None:
    docx_bytes = make_docx_bytes("Hello from Word")
    epub_bytes = make_epub_bytes("电子书标题", "Hello from EPUB")
    txt_bytes = make_txt_bytes("第一行\n第二行\n第三行")
    pdf_bytes = make_pdf_bytes("Hello from PDF")

    word_to_epub_response = client.post(
        "/api/tools/word-to-epub/upload",
        data={"params": json.dumps({"title": "Word EPUB"})},
        files={
            "file": (
                "demo.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert word_to_epub_response.status_code == 200
    assert "application/epub+zip" in word_to_epub_response.headers["content-type"]

    epub_to_word_response = client.post(
        "/api/tools/epub-to-word/upload",
        data={"params": json.dumps({"title": "EPUB Word"})},
        files={"file": ("demo.epub", epub_bytes, "application/epub+zip")},
    )
    assert epub_to_word_response.status_code == 200
    assert "wordprocessingml.document" in epub_to_word_response.headers["content-type"]

    epub_to_txt_response = client.post(
        "/api/tools/epub-to-txt/upload",
        data={"params": "{}"},
        files={"file": ("demo.epub", epub_bytes, "application/epub+zip")},
    )
    assert epub_to_txt_response.status_code == 200
    assert epub_to_txt_response.headers["content-type"].startswith("text/plain")

    word_to_html_response = client.post(
        "/api/tools/word-to-html/upload",
        data={"params": json.dumps({"title": "Word HTML"})},
        files={
            "file": (
                "demo.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert word_to_html_response.status_code == 200
    assert word_to_html_response.headers["content-type"].startswith("text/html")

    pdf_to_html_response = client.post(
        "/api/tools/pdf-to-html/upload",
        data={"params": json.dumps({"title": "PDF HTML"})},
        files={"file": ("demo.pdf", pdf_bytes, "application/pdf")},
    )
    assert pdf_to_html_response.status_code == 200
    assert pdf_to_html_response.headers["content-type"].startswith("text/html")

    txt_to_html_response = client.post(
        "/api/tools/txt-to-html/upload",
        data={"params": json.dumps({"title": "TXT HTML"})},
        files={"file": ("demo.txt", txt_bytes, "text/plain")},
    )
    assert txt_to_html_response.status_code == 200
    assert txt_to_html_response.headers["content-type"].startswith("text/html")


def test_office_tabular_and_ppt_conversion_tools_upload() -> None:
    xlsx_bytes = make_xlsx_bytes([["姓名", "分数"], ["张三", 95], ["李四", 88]], extra_sheet=True)
    csv_bytes = make_csv_bytes([["姓名", "城市"], ["张三", "上海"], ["李四", "深圳"]])
    pptx_bytes = make_pptx_bytes([["第一段内容", "第二段内容"], ["结论页", "谢谢观看"]])

    excel_to_csv_response = client.post(
        "/api/tools/excel-to-csv/upload",
        data={"params": "{}"},
        files={"file": ("demo.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert excel_to_csv_response.status_code == 200
    assert excel_to_csv_response.headers["content-type"].startswith("application/zip")

    csv_to_excel_response = client.post(
        "/api/tools/csv-to-excel/upload",
        data={"params": json.dumps({"sheet_name": "导入结果"})},
        files={"file": ("demo.csv", csv_bytes, "text/csv")},
    )
    assert csv_to_excel_response.status_code == 200
    assert "spreadsheetml.sheet" in csv_to_excel_response.headers["content-type"]

    excel_to_html_response = client.post(
        "/api/tools/excel-to-html/upload",
        data={"params": json.dumps({"title": "Excel 表格页面"})},
        files={"file": ("demo.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert excel_to_html_response.status_code == 200
    assert excel_to_html_response.headers["content-type"].startswith("text/html")

    csv_to_html_response = client.post(
        "/api/tools/csv-to-html/upload",
        data={"params": json.dumps({"title": "CSV 表格页面"})},
        files={"file": ("demo.csv", csv_bytes, "text/csv")},
    )
    assert csv_to_html_response.status_code == 200
    assert csv_to_html_response.headers["content-type"].startswith("text/html")

    ppt_to_txt_response = client.post(
        "/api/tools/ppt-to-txt/upload",
        data={"params": "{}"},
        files={"file": ("demo.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert ppt_to_txt_response.status_code == 200
    assert ppt_to_txt_response.headers["content-type"].startswith("text/plain")

    ppt_to_html_response = client.post(
        "/api/tools/ppt-to-html/upload",
        data={"params": json.dumps({"title": "PPT 页面"})},
        files={"file": ("demo.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert ppt_to_html_response.status_code == 200
    assert ppt_to_html_response.headers["content-type"].startswith("text/html")


def test_more_office_to_document_conversion_tools_upload() -> None:
    xlsx_bytes = make_xlsx_bytes([["姓名", "分数"], ["张三", 95], ["李四", 88]], extra_sheet=True)
    csv_bytes = make_csv_bytes([["姓名", "城市"], ["张三", "上海"], ["李四", "深圳"]])
    pptx_bytes = make_pptx_bytes([["第一段内容", "第二段内容"], ["结论页", "谢谢观看"]])

    ppt_to_word_response = client.post(
        "/api/tools/ppt-to-word/upload",
        data={"params": json.dumps({"title": "PPT 文档"})},
        files={"file": ("demo.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert ppt_to_word_response.status_code == 200
    assert "wordprocessingml.document" in ppt_to_word_response.headers["content-type"]

    ppt_to_pdf_response = client.post(
        "/api/tools/ppt-to-pdf/upload",
        data={"params": json.dumps({"title": "PPT PDF"})},
        files={"file": ("demo.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert ppt_to_pdf_response.status_code == 200
    assert ppt_to_pdf_response.headers["content-type"].startswith("application/pdf")

    excel_to_txt_response = client.post(
        "/api/tools/excel-to-txt/upload",
        data={"params": "{}"},
        files={"file": ("demo.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert excel_to_txt_response.status_code == 200
    assert excel_to_txt_response.headers["content-type"].startswith("text/plain")

    excel_to_pdf_response = client.post(
        "/api/tools/excel-to-pdf/upload",
        data={"params": json.dumps({"title": "Excel PDF"})},
        files={"file": ("demo.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert excel_to_pdf_response.status_code == 200
    assert excel_to_pdf_response.headers["content-type"].startswith("application/pdf")

    csv_to_txt_response = client.post(
        "/api/tools/csv-to-txt/upload",
        data={"params": "{}"},
        files={"file": ("demo.csv", csv_bytes, "text/csv")},
    )
    assert csv_to_txt_response.status_code == 200
    assert csv_to_txt_response.headers["content-type"].startswith("text/plain")

    csv_to_pdf_response = client.post(
        "/api/tools/csv-to-pdf/upload",
        data={"params": json.dumps({"title": "CSV PDF"})},
        files={"file": ("demo.csv", csv_bytes, "text/csv")},
    )
    assert csv_to_pdf_response.status_code == 200
    assert csv_to_pdf_response.headers["content-type"].startswith("application/pdf")


def test_even_more_office_to_document_conversion_tools_upload() -> None:
    xlsx_bytes = make_xlsx_bytes([["姓名", "分数"], ["张三", 95], ["李四", 88]], extra_sheet=True)
    csv_bytes = make_csv_bytes([["姓名", "城市"], ["张三", "上海"], ["李四", "深圳"]])
    pptx_bytes = make_pptx_bytes([["第一段内容", "第二段内容"], ["结论页", "谢谢观看"]])

    excel_to_word_response = client.post(
        "/api/tools/excel-to-word/upload",
        data={"params": json.dumps({"title": "Excel 文档"})},
        files={"file": ("demo.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert excel_to_word_response.status_code == 200
    assert "wordprocessingml.document" in excel_to_word_response.headers["content-type"]

    csv_to_word_response = client.post(
        "/api/tools/csv-to-word/upload",
        data={"params": json.dumps({"title": "CSV 文档"})},
        files={"file": ("demo.csv", csv_bytes, "text/csv")},
    )
    assert csv_to_word_response.status_code == 200
    assert "wordprocessingml.document" in csv_to_word_response.headers["content-type"]

    excel_to_epub_response = client.post(
        "/api/tools/excel-to-epub/upload",
        data={"params": json.dumps({"title": "Excel EPUB"})},
        files={"file": ("demo.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert excel_to_epub_response.status_code == 200
    assert "application/epub+zip" in excel_to_epub_response.headers["content-type"]

    csv_to_epub_response = client.post(
        "/api/tools/csv-to-epub/upload",
        data={"params": json.dumps({"title": "CSV EPUB"})},
        files={"file": ("demo.csv", csv_bytes, "text/csv")},
    )
    assert csv_to_epub_response.status_code == 200
    assert "application/epub+zip" in csv_to_epub_response.headers["content-type"]

    ppt_to_epub_response = client.post(
        "/api/tools/ppt-to-epub/upload",
        data={"params": json.dumps({"title": "PPT EPUB"})},
        files={"file": ("demo.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert ppt_to_epub_response.status_code == 200
    assert "application/epub+zip" in ppt_to_epub_response.headers["content-type"]


def test_generated_image_tools_execute() -> None:
    art_qr_response = client.post(
        "/api/tools/art-qr-generator/execute",
        json={
            "text": "https://qixiang.tools",
            "params": {
                "style": "rounded",
                "color_mode": "gradient",
                "fill_color": "#1d4ed8",
                "accent_color": "#14b8a6",
                "back_color": "#ffffff",
            },
        },
    )
    assert art_qr_response.status_code == 200
    assert art_qr_response.headers["content-type"].startswith("image/png")

    signature_response = client.post(
        "/api/tools/handwritten-signature/execute",
        json={"text": "张三", "params": {"color": "#111111", "font_size": 120, "transparent_bg": True}},
    )
    assert signature_response.status_code == 200
    assert signature_response.headers["content-type"].startswith("image/png")

    svg_response = client.post(
        "/api/tools/svg-placeholder-generator/execute",
        json={
            "text": "Demo Banner",
            "params": {"width": 800, "height": 450, "background_color": "#dbeafe", "text_color": "#1e3a8a"},
        },
    )
    assert svg_response.status_code == 200
    assert "image/svg+xml" in svg_response.headers["content-type"]

    stamp_response = client.post(
        "/api/tools/stamp-generator/execute",
        json={"text": "琦湘工具集合", "params": {"center_text": "专用章", "subtitle": "2026", "color": "#d40000"}},
    )
    assert stamp_response.status_code == 200
    assert stamp_response.headers["content-type"].startswith("image/png")

    certificate_response = client.post(
        "/api/tools/funny-certificate-generator/execute",
        json={
            "text": "张三\n李四",
            "params": {"title": "最佳摸鱼奖", "reason": "因摸鱼姿势过于标准", "signer": "琦湘评审团"},
        },
    )
    assert certificate_response.status_code == 200
    assert "attachment" in certificate_response.headers.get("content-disposition", "")


def test_flag_avatar_and_image_byte_expander_tools_upload() -> None:
    png_bytes = make_png_bytes((64, 128, 255), (120, 120))

    flag_response = client.post(
        "/api/tools/flag-avatar-generator/upload",
        data={"params": json.dumps({"flag_code": "cn", "opacity": 45})},
        files={"file": ("avatar.png", png_bytes, "image/png")},
    )
    assert flag_response.status_code == 200
    assert flag_response.headers["content-type"].startswith("image/png")

    target_kb = 32
    expand_response = client.post(
        "/api/tools/image-upscaler/upload",
        data={"params": json.dumps({"mode": "byte_size", "target_kb": target_kb})},
        files={"file": ("avatar.png", png_bytes, "image/png")},
    )
    assert expand_response.status_code == 200
    assert expand_response.headers["content-type"].startswith("image/png")
    assert len(expand_response.content) == target_kb * 1024
    expanded_image = Image.open(io.BytesIO(expand_response.content))
    assert expanded_image.size == (120, 120)

    resize_response = client.post(
        "/api/tools/image-upscaler/upload",
        data={"params": json.dumps({"mode": "dimensions", "scale": 3})},
        files={"file": ("avatar.png", png_bytes, "image/png")},
    )
    assert resize_response.status_code == 200
    assert resize_response.headers["content-type"].startswith("image/png")
    resized_image = Image.open(io.BytesIO(resize_response.content))
    assert resized_image.size == (360, 360)


def test_image_text_recognizer_upload() -> None:
    stream = io.BytesIO()
    image = Image.new("RGB", (360, 120), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.load_default(size=30)
    except TypeError:
        font = ImageFont.load_default()
    draw.text((24, 34), "Hello 123", fill="black", font=font)
    image.save(stream, format="PNG")

    response = client.post(
        "/api/tools/image-text-recognizer/upload",
        data={"params": "{}"},
        files={"file": ("ocr.png", stream.getvalue(), "image/png")},
    )
    assert response.status_code == 200
    result = response.json()["data"]["result"]
    normalized = result.replace(" ", "").lower()
    assert "hello" in normalized or "123" in normalized


def test_image_background_remover_upload() -> None:
    stream = io.BytesIO()
    image = Image.new("RGB", (140, 140), "white")
    drawer = ImageDraw.Draw(image)
    drawer.rounded_rectangle((28, 24, 112, 116), radius=18, fill=(29, 120, 255))
    drawer.rectangle((60, 48, 80, 92), fill=(255, 255, 255))
    image.save(stream, format="PNG")

    response = client.post(
        "/api/tools/image-background-remover/upload",
        data={"params": json.dumps({"strength": "standard", "trim_border": "no"})},
        files={"file": ("cutout.png", stream.getvalue(), "image/png")},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("image/png")

    result_image = Image.open(io.BytesIO(response.content)).convert("RGBA")
    assert result_image.getpixel((0, 0))[3] == 0
    body = result_image.getpixel((50, 60))
    assert body[3] >= 230


def test_image_background_remover_smart_checkerboard_upload() -> None:
    stream = io.BytesIO()
    image = Image.new("RGB", (180, 140), (236, 236, 236))
    drawer = ImageDraw.Draw(image)
    for x in range(0, 180, 18):
        drawer.rectangle((x, 0, x + 8, 140), fill=(210, 226, 244))
    drawer.ellipse((48, 22, 136, 118), fill=(34, 126, 255))
    drawer.rectangle((84, 48, 100, 92), fill=(255, 255, 255))
    image.save(stream, format="PNG")

    response = client.post(
        "/api/tools/image-background-remover/upload",
        data={"params": json.dumps({"mode": "smart", "strength": "standard", "output_background": "checkerboard", "trim_border": "no"})},
        files={"file": ("complex-cutout.png", stream.getvalue(), "image/png")},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("image/png")

    result_image = Image.open(io.BytesIO(response.content)).convert("RGBA")
    corner = result_image.getpixel((0, 0))
    body = result_image.getpixel((90, 70))
    assert corner[:3] != (255, 255, 255) or corner[3] == 255
    assert body[2] >= 200


def test_image_background_remover_smart_transparent_keeps_inner_light_details() -> None:
    stream = io.BytesIO()
    image = Image.new("RGB", (240, 180), "white")
    drawer = ImageDraw.Draw(image)
    drawer.ellipse((55, 30, 185, 155), fill=(34, 126, 255))
    drawer.rectangle((98, 70, 132, 110), fill=(255, 255, 255))
    image.save(stream, format="PNG")

    response = client.post(
        "/api/tools/image-background-remover/upload",
        data={
            "params": json.dumps(
                {
                    "mode": "smart",
                    "strength": "standard",
                    "output_background": "transparent",
                    "trim_border": "no",
                }
            )
        },
        files={"file": ("smart-transparent.png", stream.getvalue(), "image/png")},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("image/png")

    result_image = Image.open(io.BytesIO(response.content)).convert("RGBA")
    alpha_values = list(result_image.getchannel("A").getdata())
    transparent_ratio = sum(1 for alpha in alpha_values if alpha < 16) / len(alpha_values)
    assert transparent_ratio > 0.45
    assert result_image.getpixel((0, 0))[3] == 0
    assert result_image.getpixel((120, 90))[3] >= 230


def test_image_background_remover_smart_cleans_dark_checkerboard_logo_background() -> None:
    stream = io.BytesIO()
    image = Image.new("RGB", (260, 220), (18, 18, 22))
    drawer = ImageDraw.Draw(image)
    for y in range(0, 220, 16):
        for x in range(0, 260, 16):
            if (x // 16 + y // 16) % 2 == 0:
                drawer.rectangle((x, y, x + 15, y + 15), fill=(34, 36, 42))

    glow = Image.new("RGBA", image.size, (0, 0, 0, 0))
    glow_drawer = ImageDraw.Draw(glow)
    glow_drawer.ellipse((40, 28, 220, 204), fill=(78, 172, 255, 150))
    glow = glow.filter(ImageFilter.GaussianBlur(10))
    image = Image.alpha_composite(image.convert("RGBA"), glow).convert("RGB")
    drawer = ImageDraw.Draw(image)
    drawer.ellipse((54, 42, 206, 190), fill=(17, 47, 118), outline=(91, 185, 255), width=8)
    drawer.line((86, 70, 174, 164), fill=(110, 232, 255), width=18)
    drawer.line((174, 70, 86, 164), fill=(110, 232, 255), width=18)
    drawer.line((86, 70, 174, 164), fill=(255, 255, 255), width=5)
    drawer.line((174, 70, 86, 164), fill=(255, 255, 255), width=5)
    image.save(stream, format="PNG")

    response = client.post(
        "/api/tools/image-background-remover/upload",
        data={
            "params": json.dumps(
                {
                    "mode": "smart",
                    "strength": "strong",
                    "output_background": "transparent",
                    "trim_border": "no",
                }
            )
        },
        files={"file": ("dark-logo.png", stream.getvalue(), "image/png")},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("image/png")

    result_image = Image.open(io.BytesIO(response.content)).convert("RGBA")
    alpha_values = list(result_image.getchannel("A").getdata())
    transparent_ratio = sum(1 for alpha in alpha_values if alpha < 16) / len(alpha_values)
    assert transparent_ratio > 0.28
    assert result_image.getpixel((6, 6))[3] == 0
    assert result_image.getpixel((130, 112))[3] >= 230
    assert result_image.getpixel((55, 110))[3] >= 80


def test_image_background_remover_preserves_true_transparent_logo_png() -> None:
    stream = io.BytesIO()
    image = Image.new("RGBA", (180, 180), (180, 196, 210, 0))
    glow = Image.new("RGBA", image.size, (0, 0, 0, 0))
    glow_drawer = ImageDraw.Draw(glow)
    glow_drawer.ellipse((26, 26, 154, 154), fill=(72, 178, 255, 90))
    glow = glow.filter(ImageFilter.GaussianBlur(8))
    image = Image.alpha_composite(image, glow)
    drawer = ImageDraw.Draw(image)
    drawer.ellipse((42, 42, 138, 138), fill=(12, 80, 190, 255), outline=(116, 226, 255, 255), width=7)
    drawer.line((68, 70, 112, 116), fill=(120, 238, 255, 255), width=14)
    drawer.line((112, 70, 68, 116), fill=(120, 238, 255, 255), width=14)
    image.save(stream, format="PNG")

    response = client.post(
        "/api/tools/image-background-remover/upload",
        data={
            "params": json.dumps(
                {
                    "mode": "smart",
                    "strength": "strong",
                    "output_background": "transparent",
                    "trim_border": "no",
                }
            )
        },
        files={"file": ("true-transparent-logo.png", stream.getvalue(), "image/png")},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("image/png")

    result_image = Image.open(io.BytesIO(response.content)).convert("RGBA")
    assert result_image.size == image.size
    assert result_image.getpixel((0, 0))[3] == 0
    assert result_image.getpixel((0, 0))[:3] == (0, 0, 0)
    assert result_image.getpixel((90, 90))[3] == 255
    original_soft_alpha = image.getpixel((32, 90))[3]
    result_soft_alpha = result_image.getpixel((32, 90))[3]
    assert abs(result_soft_alpha - original_soft_alpha) <= 2


def test_uuid_generator_execute() -> None:
    response = client.post(
        "/api/tools/uuid-generator/execute",
        json={"text": "", "params": {"version": "v4", "count": 5}},
    )
    assert response.status_code == 200
    values = response.json()["data"]["result"].splitlines()
    assert len(values) == 5
    for value in values:
        assert str(uuid.UUID(value)) == value


def test_json_url_params_execute() -> None:
    response = client.post(
        "/api/tools/json-url-params/execute",
        json={"text": '{"name":"alice","tags":["x","y"]}', "params": {"action": "json_to_params"}},
    )
    assert response.status_code == 200
    parsed = parse_qs(response.json()["data"]["result"])
    assert parsed == {"name": ["alice"], "tags": ["x", "y"]}


def test_json_xml_converter_execute() -> None:
    response = client.post(
        "/api/tools/json-xml-converter/execute",
        json={"text": '{"name":"Alice","age":18}', "params": {"action": "json_to_xml"}},
    )
    assert response.status_code == 200
    result = response.json()["data"]["result"]
    assert "<name>Alice</name>" in result
    assert "<age type=\"number\">18</age>" in result


def test_json_to_typescript_execute() -> None:
    response = client.post(
        "/api/tools/json-to-typescript/execute",
        json={"text": '{"name":"Alice","age":18}', "params": {"root_name": "UserProfile"}},
    )
    assert response.status_code == 200
    result = response.json()["data"]["result"]
    assert "export interface UserProfile" in result
    assert "name: string;" in result
    assert "age: number;" in result


def test_jsonpath_parser_execute() -> None:
    response = client.post(
        "/api/tools/jsonpath-parser/execute",
        json={
            "text": '{"data":{"list":[{"name":"Alice"},{"name":"Bob"}]}}',
            "params": {"path": "$.data.list[*].name"},
        },
    )
    assert response.status_code == 200
    payload = json.loads(response.json()["data"]["result"])
    assert payload["match_count"] == 2
    assert payload["matches"] == ["Alice", "Bob"]


def test_xml_formatter_execute() -> None:
    response = client.post(
        "/api/tools/xml-formatter/execute",
        json={"text": "<root><item>1</item><item>2</item></root>", "params": {"action": "format"}},
    )
    assert response.status_code == 200
    assert "\n" in response.json()["data"]["result"]


def test_directory_tree_generator_execute() -> None:
    response = client.post(
        "/api/tools/directory-tree-generator/execute",
        json={"text": "src/main.py\nsrc/utils/helpers.py\nREADME.md", "params": {}},
    )
    assert response.status_code == 200
    result = response.json()["data"]["result"]
    assert "src" in result
    assert "main.py" in result
    assert "README.md" in result


def test_color_and_css_tools_execute() -> None:
    color_response = client.post(
        "/api/tools/color-converter/execute",
        json={"text": "#ff8800", "params": {"to_format": "rgb"}},
    )
    assert color_response.status_code == 200
    assert color_response.json()["data"]["result"] == "rgb(255, 136, 0)"

    unit_response = client.post(
        "/api/tools/css-unit-converter/execute",
        json={
            "text": "32",
            "params": {
                "from_unit": "px",
                "to_unit": "rem",
                "root_font_size": "16",
                "screen_width": "375",
            },
        },
    )
    assert unit_response.status_code == 200
    assert unit_response.json()["data"]["result"] == "2rem"

    gradient_response = client.post(
        "/api/tools/css-gradient-generator/execute",
        json={
            "text": "",
            "params": {
                "gradient_type": "linear",
                "direction": "135deg",
                "start_color": "#111111",
                "end_color": "#eeeeee",
            },
        },
    )
    assert gradient_response.status_code == 200
    assert "linear-gradient(135deg, #111111, #eeeeee)" in gradient_response.json()["data"]["result"]


def test_json_yaml_converter_execute() -> None:
    response = client.post(
        "/api/tools/json-yaml-converter/execute",
        json={"text": '{"name":"alice","age":18}', "params": {"action": "json_to_yaml"}},
    )
    assert response.status_code == 200
    result = response.json()["data"]["result"]
    assert "name: alice" in result
    assert "age: 18" in result


def test_html_javascript_css_formatters_execute() -> None:
    html_response = client.post(
        "/api/tools/html-formatter/execute",
        json={"text": "<div><span>Hello</span></div>", "params": {"action": "format"}},
    )
    assert html_response.status_code == 200
    assert "\n" in html_response.json()["data"]["result"]

    js_response = client.post(
        "/api/tools/javascript-formatter/execute",
        json={"text": "function test(){return 1+2;}", "params": {"action": "minify"}},
    )
    assert js_response.status_code == 200
    assert "function test(){return 1+2;}" == js_response.json()["data"]["result"]

    css_response = client.post(
        "/api/tools/css-formatter/execute",
        json={"text": "body { color: red; }", "params": {"action": "minify"}},
    )
    assert css_response.status_code == 200
    assert css_response.json()["data"]["result"] == "body{color:red}"


def test_crontab_next_run_execute() -> None:
    response = client.post(
        "/api/tools/crontab-next-run/execute",
        json={"text": "*/5 * * * *", "params": {"count": 5}},
    )
    assert response.status_code == 200
    result = response.json()["data"]["result"]
    assert "未来执行时间:" in result
    assert "1." in result


def test_base_and_byte_converters_execute() -> None:
    base_response = client.post(
        "/api/tools/base-converter/execute",
        json={"text": "255", "params": {"from_base": 10, "to_base": 16}},
    )
    assert base_response.status_code == 200
    assert base_response.json()["data"]["result"] == "FF"

    byte_response = client.post(
        "/api/tools/byte-converter/execute",
        json={"text": "1048576", "params": {"from_unit": "B", "to_unit": "MiB"}},
    )
    assert byte_response.status_code == 200
    assert byte_response.json()["data"]["result"] == "1 MiB"


def test_file_hash_upload() -> None:
    content = b"hello hash"
    response = client.post(
        "/api/tools/file-hash/upload",
        data={"params": json.dumps({"algorithm": "sha256"})},
        files={"file": ("hash.txt", content, "text/plain")},
    )
    assert response.status_code == 200
    assert response.json()["data"]["result"] == hashlib.sha256(content).hexdigest()


def test_text_sorter_execute() -> None:
    response = client.post(
        "/api/tools/text-sorter/execute",
        json={
            "text": "c\na\nb\na",
            "params": {"mode": "text", "order": "asc", "unique": True},
        },
    )
    assert response.status_code == 200
    assert response.json()["data"]["result"] == "a\nb\nc"


def test_text_statistics_execute() -> None:
    response = client.post(
        "/api/tools/text-statistics/execute",
        json={"text": "hi\nhello world", "params": {}},
    )
    assert response.status_code == 200
    result = response.json()["data"]["result"]
    assert "总字符数：14" in result
    assert "总行数：2" in result
    assert "非空行数：2" in result
    assert "单词数量：3" in result


def test_html_text_extractor_execute() -> None:
    response = client.post(
        "/api/tools/html-text-extractor/execute",
        json={
            "text": "<div>Hello <b>World</b><script>alert(1)</script><style>.x{}</style></div>",
            "params": {"join_mode": "single_line"},
        },
    )
    assert response.status_code == 200
    assert response.json()["data"]["result"] == "Hello World"


def test_text_transformers_execute() -> None:
    english_response = client.post(
        "/api/tools/english-case-converter/execute",
        json={"text": "hello world. this is codex", "params": {"action": "sentence"}},
    )
    assert english_response.status_code == 200
    assert english_response.json()["data"]["result"] == "Hello world. This is codex"

    pinyin_response = client.post(
        "/api/tools/pinyin-converter/execute",
        json={"text": "重庆银行", "params": {"style": "first_letter", "separator": "empty"}},
    )
    assert pinyin_response.status_code == 200
    assert pinyin_response.json()["data"]["result"] == "cqyh"

    convert_response = client.post(
        "/api/tools/traditional-simplified-converter/execute",
        json={"text": "简体中文与开发工具", "params": {"action": "simplified_to_traditional"}},
    )
    assert convert_response.status_code == 200
    assert "簡體中文" in convert_response.json()["data"]["result"]

    halfwidth_response = client.post(
        "/api/tools/fullwidth-halfwidth-converter/execute",
        json={"text": "ABC 123", "params": {"action": "to_fullwidth"}},
    )
    assert halfwidth_response.status_code == 200
    assert halfwidth_response.json()["data"]["result"] == "ＡＢＣ　１２３"

    mars_response = client.post(
        "/api/tools/mars-text-converter/execute",
        json={"text": "你好123", "params": {"action": "encode"}},
    )
    assert mars_response.status_code == 200
    mars_text = mars_response.json()["data"]["result"]
    assert "伱" in mars_text
    assert "１２３" in mars_text

    mars_decode_response = client.post(
        "/api/tools/mars-text-converter/execute",
        json={"text": mars_text, "params": {"action": "decode"}},
    )
    assert mars_decode_response.status_code == 200
    assert mars_decode_response.json()["data"]["result"] == "你好123"


def test_text_diff_regex_and_rule_tools_execute() -> None:
    diff_response = client.post(
        "/api/tools/text-diff-comparer/execute",
        json={
            "text": "a\nb\nc",
            "params": {"compare_text": "a\nx\nc", "context_lines": 1},
        },
    )
    assert diff_response.status_code == 200
    diff_result = diff_response.json()["data"]["result"]
    assert "-b" in diff_result
    assert "+x" in diff_result

    regex_response = client.post(
        "/api/tools/regex-replacer/execute",
        json={
            "text": "abc123def456",
            "params": {"pattern": "\\d+", "replacement": "#", "flags": "", "count": 1},
        },
    )
    assert regex_response.status_code == 200
    assert regex_response.json()["data"]["result"] == "abc#def456"

    rule_response = client.post(
        "/api/tools/text-rule-generator/execute",
        json={
            "text": "apple\nbanana\napple",
            "params": {"mode": "exact_line", "ignore_case": True},
        },
    )
    assert rule_response.status_code == 200
    assert rule_response.json()["data"]["result"] == "(?i)^(?:apple|banana)$"


def test_text_extractors_and_number_tools_execute() -> None:
    email_response = client.post(
        "/api/tools/email-extractor/execute",
        json={
            "text": "联系 a@test.com，再抄送 a@test.com 和 b@example.org",
            "params": {"unique": True},
        },
    )
    assert email_response.status_code == 200
    assert email_response.json()["data"]["result"] == "a@test.com\nb@example.org"

    mobile_response = client.post(
        "/api/tools/mobile-extractor/execute",
        json={
            "text": "13800000000，备用 +86 139 1111 2222",
            "params": {"unique": True},
        },
    )
    assert mobile_response.status_code == 200
    assert mobile_response.json()["data"]["result"] == "13800000000\n13911112222"

    link_response = client.post(
        "/api/tools/link-extractor/execute",
        json={
            "text": "https://example.com/test?x=1 还有 www.openai.com.",
            "params": {"unique": True},
        },
    )
    assert link_response.status_code == 200
    assert link_response.json()["data"]["result"] == "https://example.com/test?x=1\nwww.openai.com"

    number_response = client.post(
        "/api/tools/line-number-adder/execute",
        json={
            "text": "Alpha\n\nBeta",
            "params": {"start_number": 1, "style": "wrap_paren", "keep_empty": False},
        },
    )
    assert number_response.status_code == 200
    assert number_response.json()["data"]["result"] == "(1) Alpha\n\n(2) Beta"


def test_text_generation_and_layout_tools_execute() -> None:
    name_response = client.post(
        "/api/tools/name-generator/execute",
        json={"text": "", "params": {"count": 5, "gender": "mixed"}},
    )
    assert name_response.status_code == 200
    names = name_response.json()["data"]["result"].splitlines()
    assert len(names) == 5
    assert all(len(name) >= 2 for name in names)

    camel_response = client.post(
        "/api/tools/camel-snake-converter/execute",
        json={"text": "hello_world_value", "params": {"action": "snake_to_camel"}},
    )
    assert camel_response.status_code == 200
    assert camel_response.json()["data"]["result"] == "helloWorldValue"

    vertical_response = client.post(
        "/api/tools/vertical-text-tool/execute",
        json={"text": "天地", "params": {"blank_line_between_paragraphs": False}},
    )
    assert vertical_response.status_code == 200
    assert vertical_response.json()["data"]["result"] == "天\n地"

    copybook_response = client.post(
        "/api/tools/copybook-generator/execute",
        json={"text": "天", "params": {"repeat_count": 3}},
    )
    assert copybook_response.status_code == 200
    assert copybook_response.json()["data"]["result"] == "天　天　天"


def test_superscript_subscript_and_word_group_tools_execute() -> None:
    superscript_response = client.post(
        "/api/tools/superscript-phone-generator/execute",
        json={"text": "+86 138", "params": {}},
    )
    assert superscript_response.status_code == 200
    assert superscript_response.json()["data"]["result"] == "⁺⁸⁶ ¹³⁸"

    subscript_response = client.post(
        "/api/tools/subscript-phone-generator/execute",
        json={"text": "138", "params": {}},
    )
    assert subscript_response.status_code == 200
    assert subscript_response.json()["data"]["result"] == "₁₃₈"

    group_response = client.post(
        "/api/tools/word-group-builder/execute",
        json={"text": "汉", "params": {"level": "high", "limit": 5}},
    )
    assert group_response.status_code == 200
    group_result = group_response.json()["data"]["result"]
    assert "汉语" in group_result or "武汉" in group_result


def test_wubi_and_stroke_tools_execute() -> None:
    wubi_response = client.post(
        "/api/tools/wubi-converter/execute",
        json={"text": "你好", "params": {"mode": "both"}},
    )
    assert wubi_response.status_code == 200
    wubi_result = wubi_response.json()["data"]["result"]
    assert "你：wqiy" in wubi_result
    assert "词组编码：" in wubi_result
    assert "wqvb" in wubi_result

    stroke_response = client.post(
        "/api/tools/stroke-query/execute",
        json={"text": "你好", "params": {}},
    )
    assert stroke_response.status_code == 200
    stroke_result = stroke_response.json()["data"]["result"]
    assert "你：7 画" in stroke_result
    assert "好：6 画" in stroke_result


def test_idiom_and_xiehouyu_tools_execute() -> None:
    idiom_response = client.post(
        "/api/tools/idiom-cloze/execute",
        json={"text": "一?一?", "params": {"limit": 5}},
    )
    assert idiom_response.status_code == 200
    idiom_result = idiom_response.json()["data"]["result"]
    assert "一心一意" in idiom_result

    xiehouyu_response = client.post(
        "/api/tools/xiehouyu-search/execute",
        json={"text": "八仙过海", "params": {"limit": 5}},
    )
    assert xiehouyu_response.status_code == 200
    xiehouyu_result = xiehouyu_response.json()["data"]["result"]
    assert "各显神通" in xiehouyu_result


def test_text_cleanup_and_conversion_tools_execute() -> None:
    reverse_response = client.post(
        "/api/tools/text-reverser/execute",
        json={"text": "abc\n123", "params": {"mode": "chars_per_line"}},
    )
    assert reverse_response.status_code == 200
    assert reverse_response.json()["data"]["result"] == "cba\n321"

    whitespace_response = client.post(
        "/api/tools/whitespace-cleaner/execute",
        json={
            "text": "  hello   world  \n\n  second   line ",
            "params": {
                "trim_each_line": True,
                "remove_empty_lines": True,
                "collapse_spaces": True,
            },
        },
    )
    assert whitespace_response.status_code == 200
    assert whitespace_response.json()["data"]["result"] == "hello world\nsecond line"

    format_cleaner_response = client.post(
        "/api/tools/text-format-cleaner/execute",
        json={"text": "  A\tB\u200b  \n\n  C\r\nD  ", "params": {}},
    )
    assert format_cleaner_response.status_code == 200
    assert format_cleaner_response.json()["data"]["result"] == "A B C D"

    keep_lines_response = client.post(
        "/api/tools/text-format-cleaner/execute",
        json={"text": "  A\tB\u200b  \n\n  C   D  ", "params": {"mode": "keep_lines"}},
    )
    assert keep_lines_response.status_code == 200
    assert keep_lines_response.json()["data"]["result"] == "A B\nC D"

    invisible_response = client.post(
        "/api/tools/invisible-control-chars/execute",
        json={"text": "", "params": {}},
    )
    assert invisible_response.status_code == 200
    invisible_result = invisible_response.json()["data"]["result"]
    assert "U+2028" in invisible_result
    assert "LINE SEPARATOR" in invisible_result
    assert "所有 Raw Char" in invisible_result
    assert "以上是所有的不可见字符" in invisible_result

    raw_only_response = client.post(
        "/api/tools/invisible-control-chars/execute",
        json={"text": "", "params": {"output_mode": "raw_only", "joiner": "comma"}},
    )
    assert raw_only_response.status_code == 200
    raw_only_result = raw_only_response.json()["data"]["result"]
    assert "、" in raw_only_result
    assert "U+2028" not in raw_only_result
    assert raw_only_result.endswith("以上是所有的不可见字符")

    punctuation_response = client.post(
        "/api/tools/punctuation-converter/execute",
        json={"text": "你好，世界！", "params": {"action": "cn_to_en"}},
    )
    assert punctuation_response.status_code == 200
    assert punctuation_response.json()["data"]["result"] == "你好,世界!"


def test_duplicate_split_and_prefix_tools_execute() -> None:
    duplicate_response = client.post(
        "/api/tools/duplicate-line-counter/execute",
        json={
            "text": "apple\nbanana\napple\napple\npear",
            "params": {"include_single": False, "sort_by": "count_desc"},
        },
    )
    assert duplicate_response.status_code == 200
    assert duplicate_response.json()["data"]["result"] == "apple  x 3"

    split_response = client.post(
        "/api/tools/text-split-joiner/execute",
        json={
            "text": "a,b, c ,,d",
            "params": {
                "action": "split_to_lines",
                "input_separator": "comma",
                "custom_input_separator": "",
                "output_separator": "comma",
                "custom_output_separator": "",
                "trim_items": True,
                "remove_empty_items": True,
            },
        },
    )
    assert split_response.status_code == 200
    assert split_response.json()["data"]["result"] == "a\nb\nc\nd"

    join_response = client.post(
        "/api/tools/text-split-joiner/execute",
        json={
            "text": "a\nb\nc",
            "params": {
                "action": "join_lines",
                "input_separator": "comma",
                "custom_input_separator": "",
                "output_separator": "pipe",
                "custom_output_separator": "",
                "trim_items": True,
                "remove_empty_items": True,
            },
        },
    )
    assert join_response.status_code == 200
    assert join_response.json()["data"]["result"] == "a|b|c"

    prefix_response = client.post(
        "/api/tools/prefix-suffix-adder/execute",
        json={
            "text": "alpha\n\nbeta",
            "params": {"prefix": '"', "suffix": '",', "skip_empty": True},
        },
    )
    assert prefix_response.status_code == 200
    assert prefix_response.json()["data"]["result"] == '"alpha",\n\n"beta",'


def test_random_phone_generator_execute() -> None:
    response = client.post(
        "/api/tools/random-phone-generator/execute",
        json={"text": "", "params": {"count": 5}},
    )
    assert response.status_code == 200
    values = response.json()["data"]["result"].splitlines()
    assert len(values) == 5
    for value in values:
        assert len(value) == 11
        assert value.isdigit()


def test_random_id_card_generator_execute() -> None:
    response = client.post(
        "/api/tools/random-id-card-generator/execute",
        json={"text": "", "params": {"gender": "female", "count": 5}},
    )
    assert response.status_code == 200
    values = response.json()["data"]["result"].splitlines()
    assert len(values) == 5
    for value in values:
        assert len(value) == 18
        assert value[:17].isdigit()
        assert value[-1] in "0123456789X"


def test_markdown_editor_execute() -> None:
    response = client.post(
        "/api/tools/markdown-editor/execute",
        json={"text": "# 标题\n\n- A\n- B", "params": {}},
    )
    assert response.status_code == 200
    result = response.json()["data"]["result"]
    assert "<h1" in result
    assert "标题" in result
    assert "<li>A</li>" in result


def test_coordinate_converter_execute() -> None:
    response = client.post(
        "/api/tools/coordinate-converter/execute",
        json={
            "text": "116.404,39.915",
            "params": {"from_system": "wgs84", "to_system": "gcj02", "precision": 6},
        },
    )
    assert response.status_code == 200
    payload = json.loads(response.json()["data"]["result"])
    assert payload["from_system"] == "wgs84"
    assert payload["to_system"] == "gcj02"
    assert payload["output"]["lng"] != 116.404
    assert payload["output"]["lat"] != 39.915


def test_json_to_xlsx_execute() -> None:
    response = client.post(
        "/api/tools/json-to-xlsx/execute",
        json={
            "text": '[{"name":"Alice","age":18},{"name":"Bob","age":20}]',
            "params": {"sheet_name": "Users"},
        },
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers.get("content-disposition", "")


def test_xlsx_to_json_upload() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Users"
    sheet.append(["name", "age"])
    sheet.append(["Alice", 18])
    sheet.append(["Bob", 20])

    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)

    response = client.post(
        "/api/tools/xlsx-to-json/upload",
        data={"params": "{}"},
        files={
            "file": (
                "users.xlsx",
                stream.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200
    payload = json.loads(response.json()["data"]["result"])
    assert payload[0]["name"] == "Alice"
    assert payload[1]["age"] == 20


def test_sqlite_viewer_upload() -> None:
    with tempfile.NamedTemporaryFile(suffix=".db") as temp_file:
        connection = sqlite3.connect(temp_file.name)
        connection.execute("CREATE TABLE users (id INTEGER PRIMARY KEY, name TEXT)")
        connection.execute("INSERT INTO users (name) VALUES ('Alice')")
        connection.execute("INSERT INTO users (name) VALUES ('Bob')")
        connection.commit()
        connection.close()

        content = Path(temp_file.name).read_bytes()

    response = client.post(
        "/api/tools/sqlite-viewer/upload",
        data={"params": json.dumps({"preview_limit": 10})},
        files={"file": ("sample.db", content, "application/octet-stream")},
    )
    assert response.status_code == 200
    payload = json.loads(response.json()["data"]["result"])
    assert payload["table_count"] == 1
    assert payload["tables"][0]["name"] == "users"
    assert payload["tables"][0]["preview"][0]["name"] == "Alice"


def test_unicode_hex_and_punycode_codecs_execute() -> None:
    unicode_response = client.post(
        "/api/tools/unicode-chinese-converter/execute",
        json={"text": "你好", "params": {"action": "to_unicode"}},
    )
    assert unicode_response.status_code == 200
    escaped = unicode_response.json()["data"]["result"]
    assert "\\u4f60" in escaped

    unicode_decode_response = client.post(
        "/api/tools/unicode-chinese-converter/execute",
        json={"text": escaped, "params": {"action": "to_text"}},
    )
    assert unicode_decode_response.status_code == 200
    assert unicode_decode_response.json()["data"]["result"] == "你好"

    hex_response = client.post(
        "/api/tools/hex-string-codec/execute",
        json={"text": "hello", "params": {"action": "text_to_hex"}},
    )
    assert hex_response.status_code == 200
    assert hex_response.json()["data"]["result"] == "68656c6c6f"

    hex_decode_response = client.post(
        "/api/tools/hex-string-codec/execute",
        json={"text": "68656c6c6f", "params": {"action": "hex_to_text"}},
    )
    assert hex_decode_response.status_code == 200
    assert hex_decode_response.json()["data"]["result"] == "hello"

    puny_response = client.post(
        "/api/tools/punycode-codec/execute",
        json={"text": "你好世界", "params": {"action": "encode"}},
    )
    assert puny_response.status_code == 200
    puny_text = puny_response.json()["data"]["result"]

    puny_decode_response = client.post(
        "/api/tools/punycode-codec/execute",
        json={"text": puny_text, "params": {"action": "decode"}},
    )
    assert puny_decode_response.status_code == 200
    assert puny_decode_response.json()["data"]["result"] == "你好世界"


def test_md5_sha_shake_and_keccak_execute() -> None:
    md5_response = client.post(
        "/api/tools/md5-decoder/execute",
        json={"text": "hello", "params": {"action": "hash", "digest_length": 32, "digest_case": "upper"}},
    )
    assert md5_response.status_code == 200
    assert md5_response.json()["data"]["result"] == hashlib.md5(b"hello").hexdigest().upper()

    sha_response = client.post(
        "/api/tools/sha-hash/execute",
        json={"text": "hello", "params": {"algorithm": "sha256", "digest_case": "lower"}},
    )
    assert sha_response.status_code == 200
    assert sha_response.json()["data"]["result"] == hashlib.sha256(b"hello").hexdigest()

    shake_response = client.post(
        "/api/tools/shake-hash/execute",
        json={"text": "hello", "params": {"algorithm": "shake_128", "digest_bytes": 16}},
    )
    assert shake_response.status_code == 200
    assert shake_response.json()["data"]["result"] == hashlib.shake_128(b"hello").hexdigest(16)

    keccak_response = client.post(
        "/api/tools/keccak-hash/execute",
        json={"text": "hello", "params": {"digest_bits": 256, "digest_case": "lower"}},
    )
    assert keccak_response.status_code == 200
    hasher = keccak.new(digest_bits=256)
    hasher.update(b"hello")
    assert keccak_response.json()["data"]["result"] == hasher.hexdigest()


def test_jwt_decoder_execute() -> None:
    header = base64.urlsafe_b64encode(b'{"alg":"HS256","typ":"JWT"}').rstrip(b"=").decode("ascii")
    payload = base64.urlsafe_b64encode(b'{"sub":"alice","role":"admin"}').rstrip(b"=").decode("ascii")
    token = f"{header}.{payload}.signature"
    response = client.post(
        "/api/tools/jwt-decoder/execute",
        json={"text": token, "params": {}},
    )
    assert response.status_code == 200
    result = json.loads(response.json()["data"]["result"])
    assert result["header"]["alg"] == "HS256"
    assert result["payload"]["sub"] == "alice"
    assert result["signature"] == "signature"


def test_morse_and_text_hidden_codecs_execute() -> None:
    morse_response = client.post(
        "/api/tools/morse-codec/execute",
        json={"text": "SOS", "params": {"action": "encode"}},
    )
    assert morse_response.status_code == 200
    assert morse_response.json()["data"]["result"] == "... --- ..."

    morse_decode_response = client.post(
        "/api/tools/morse-codec/execute",
        json={"text": "... --- ...", "params": {"action": "decode"}},
    )
    assert morse_decode_response.status_code == 200
    assert morse_decode_response.json()["data"]["result"] == "SOS"

    hide_response = client.post(
        "/api/tools/text-hidden-codec/execute",
        json={"text": "secret", "params": {"action": "hide", "cover_text": "hello"}},
    )
    assert hide_response.status_code == 200
    hidden_text = hide_response.json()["data"]["result"]
    assert hidden_text.startswith("hello")

    reveal_response = client.post(
        "/api/tools/text-hidden-codec/execute",
        json={"text": hidden_text, "params": {"action": "reveal", "cover_text": ""}},
    )
    assert reveal_response.status_code == 200
    assert reveal_response.json()["data"]["result"] == "secret"


def test_bcrypt_tool_execute() -> None:
    hash_response = client.post(
        "/api/tools/bcrypt-tool/execute",
        json={"text": "pa55word", "params": {"action": "hash", "rounds": 10, "hash_value": ""}},
    )
    assert hash_response.status_code == 200
    hashed = hash_response.json()["data"]["result"]
    assert hashed.startswith("$2")

    verify_response = client.post(
        "/api/tools/bcrypt-tool/execute",
        json={"text": "pa55word", "params": {"action": "verify", "rounds": 10, "hash_value": hashed}},
    )
    assert verify_response.status_code == 200
    assert json.loads(verify_response.json()["data"]["result"])["matched"] is True


def test_aes_des_rc4_ciphers_execute() -> None:
    aes_encrypt = client.post(
        "/api/tools/aes-cipher/execute",
        json={
            "text": "hello aes",
            "params": {
                "action": "encrypt",
                "key_text": "1234567890abcdef",
                "key_encoding": "utf8",
                "output_encoding": "base64",
            },
        },
    )
    assert aes_encrypt.status_code == 200
    aes_cipher = aes_encrypt.json()["data"]["result"]

    aes_decrypt = client.post(
        "/api/tools/aes-cipher/execute",
        json={
            "text": aes_cipher,
            "params": {
                "action": "decrypt",
                "key_text": "1234567890abcdef",
                "key_encoding": "utf8",
                "output_encoding": "base64",
            },
        },
    )
    assert aes_decrypt.status_code == 200
    assert aes_decrypt.json()["data"]["result"] == "hello aes"

    des_encrypt = client.post(
        "/api/tools/des-cipher/execute",
        json={
            "text": "hello des",
            "params": {
                "action": "encrypt",
                "key_text": "12345678",
                "key_encoding": "utf8",
                "output_encoding": "hex",
            },
        },
    )
    assert des_encrypt.status_code == 200
    des_cipher = des_encrypt.json()["data"]["result"]

    des_decrypt = client.post(
        "/api/tools/des-cipher/execute",
        json={
            "text": des_cipher,
            "params": {
                "action": "decrypt",
                "key_text": "12345678",
                "key_encoding": "utf8",
                "output_encoding": "hex",
            },
        },
    )
    assert des_decrypt.status_code == 200
    assert des_decrypt.json()["data"]["result"] == "hello des"

    rc4_encrypt = client.post(
        "/api/tools/rc4-cipher/execute",
        json={
            "text": "hello rc4",
            "params": {
                "action": "encrypt",
                "key_text": "secret-key",
                "key_encoding": "utf8",
                "output_encoding": "base64",
            },
        },
    )
    assert rc4_encrypt.status_code == 200
    rc4_cipher = rc4_encrypt.json()["data"]["result"]

    rc4_decrypt = client.post(
        "/api/tools/rc4-cipher/execute",
        json={
            "text": rc4_cipher,
            "params": {
                "action": "decrypt",
                "key_text": "secret-key",
                "key_encoding": "utf8",
                "output_encoding": "base64",
            },
        },
    )
    assert rc4_decrypt.status_code == 200
    assert rc4_decrypt.json()["data"]["result"] == "hello rc4"


def test_rsa_cipher_execute() -> None:
    key = RSA.generate(2048)
    private_key = key.export_key().decode("utf-8")
    public_key = key.publickey().export_key().decode("utf-8")

    encrypt_response = client.post(
        "/api/tools/rsa-cipher/execute",
        json={
            "text": "hello rsa",
            "params": {
                "action": "encrypt",
                "key_text": public_key,
                "output_encoding": "base64",
            },
        },
    )
    assert encrypt_response.status_code == 200
    encrypted = encrypt_response.json()["data"]["result"]

    decrypt_response = client.post(
        "/api/tools/rsa-cipher/execute",
        json={
            "text": encrypted,
            "params": {
                "action": "decrypt",
                "key_text": private_key,
                "output_encoding": "base64",
            },
        },
    )
    assert decrypt_response.status_code == 200
    assert decrypt_response.json()["data"]["result"] == "hello rsa"


def test_sm4_cipher_execute() -> None:
    encrypt_response = client.post(
        "/api/tools/sm4-cipher/execute",
        json={
            "text": "hello sm4",
            "params": {
                "action": "encrypt",
                "key_text": "1234567890abcdef",
                "key_encoding": "utf8",
                "output_encoding": "base64",
            },
        },
    )
    assert encrypt_response.status_code == 200
    encrypted = encrypt_response.json()["data"]["result"]

    decrypt_response = client.post(
        "/api/tools/sm4-cipher/execute",
        json={
            "text": encrypted,
            "params": {
                "action": "decrypt",
                "key_text": "1234567890abcdef",
                "key_encoding": "utf8",
                "output_encoding": "base64",
            },
        },
    )
    assert decrypt_response.status_code == 200
    assert decrypt_response.json()["data"]["result"] == "hello sm4"


def test_sm2_cipher_execute() -> None:
    generate_response = client.post(
        "/api/tools/sm2-cipher/execute",
        json={"text": "", "params": {"action": "generate_keypair", "public_key": "", "private_key": "", "output_encoding": "base64"}},
    )
    assert generate_response.status_code == 200
    key_text = generate_response.json()["data"]["result"]
    parts = [item.strip() for item in key_text.splitlines() if item.strip() and not item.endswith(":")]
    private_key, public_key = parts[0], parts[1]

    encrypt_response = client.post(
        "/api/tools/sm2-cipher/execute",
        json={
            "text": "hello sm2",
            "params": {
                "action": "encrypt",
                "public_key": public_key,
                "private_key": "",
                "output_encoding": "base64",
            },
        },
    )
    assert encrypt_response.status_code == 200
    encrypted = encrypt_response.json()["data"]["result"]

    decrypt_response = client.post(
        "/api/tools/sm2-cipher/execute",
        json={
            "text": encrypted,
            "params": {
                "action": "decrypt",
                "public_key": public_key,
                "private_key": private_key,
                "output_encoding": "base64",
            },
        },
    )
    assert decrypt_response.status_code == 200
    assert decrypt_response.json()["data"]["result"] == "hello sm2"


def test_md5_decoder_and_javascript_obfuscator_execute() -> None:
    md5_hash_response = client.post(
        "/api/tools/md5-decoder/execute",
        json={
            "text": "hello",
            "params": {
                "action": "hash",
                "digest_length": 16,
                "digest_case": "lower",
            },
        },
    )
    assert md5_hash_response.status_code == 200
    assert md5_hash_response.json()["data"]["result"] == hashlib.md5(b"hello").hexdigest()[8:24]

    md5_response = client.post(
        "/api/tools/md5-decoder/execute",
        json={
            "text": hashlib.md5(b"admin123").hexdigest(),
            "params": {
                "action": "decode",
                "include_common": True,
                "custom_candidates": "",
            },
        },
    )
    assert md5_response.status_code == 200
    md5_payload = json.loads(md5_response.json()["data"]["result"])
    assert md5_payload["matched"][0]["plain"] == "admin123"

    js_response = client.post(
        "/api/tools/javascript-obfuscator/execute",
        json={
            "text": "function test(){console.log('ok');}",
            "params": {
                "mode": "base64_wrapper",
                "minify_first": True,
            },
        },
    )
    assert js_response.status_code == 200
    assert "eval" in js_response.json()["data"]["result"]


def test_rabbit_cipher_execute_and_vector() -> None:
    encrypt_response = client.post(
        "/api/tools/rabbit-cipher/execute",
        json={
            "text": "hello rabbit",
            "params": {
                "action": "encrypt",
                "key_text": "1234567890abcdef",
                "key_encoding": "utf8",
                "iv_text": "12345678",
                "iv_encoding": "utf8",
                "output_encoding": "base64",
            },
        },
    )
    assert encrypt_response.status_code == 200
    encrypted = encrypt_response.json()["data"]["result"]

    decrypt_response = client.post(
        "/api/tools/rabbit-cipher/execute",
        json={
            "text": encrypted,
            "params": {
                "action": "decrypt",
                "key_text": "1234567890abcdef",
                "key_encoding": "utf8",
                "iv_text": "",
                "iv_encoding": "utf8",
                "output_encoding": "base64",
            },
        },
    )
    assert decrypt_response.status_code == 200
    assert decrypt_response.json()["data"]["result"] == "hello rabbit"

    key = bytes.fromhex("00000000000000000000000000000000")
    iv = bytes.fromhex("0000000000000000")
    stream = RabbitCipher(key=key, iv=iv).crypt(bytes(16))
    assert stream.hex().upper() == "C6A7275EF85495D87CCD5D376705B7ED"
