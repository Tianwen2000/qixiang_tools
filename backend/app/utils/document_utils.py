"""文件说明：提供 Word、PDF、PPT、Excel 等文档处理公共函数。"""

import csv
import re
from functools import lru_cache
from html import escape
from io import StringIO
from pathlib import Path
from uuid import uuid4

import fitz
from bs4 import BeautifulSoup
from docx import Document
from ebooklib import ITEM_DOCUMENT, epub
from openpyxl import Workbook, load_workbook
from PIL import Image
from pillow_heif import register_heif_opener
from pptx import Presentation
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from app.core.exceptions import AppException
from app.utils.image_utils import ensure_rgb, zip_output_files


PDF_FONT_NAME = "STSong-Light"


@lru_cache(maxsize=1)
def get_pdf_font_name() -> str:
    try:
        pdfmetrics.getFont(PDF_FONT_NAME)
    except KeyError:
        pdfmetrics.registerFont(UnicodeCIDFont(PDF_FONT_NAME))
    return PDF_FONT_NAME


def normalize_extracted_text(text: str) -> str:
    value = text.replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r"[ \t]+\n", "\n", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


@lru_cache(maxsize=1)
def register_heic_support() -> bool:
    register_heif_opener()
    return True


def read_plain_text(input_path: str | Path) -> str:
    return normalize_extracted_text(read_text_content(input_path))


def read_text_content(input_path: str | Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "gbk", "big5"):
        try:
            return Path(input_path).read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
        except Exception as exc:
            raise AppException(message=f"无法读取文本文件：{exc}", code=4002, status_code=400) from exc
    raise AppException(message="无法识别文本文件编码", code=4002, status_code=400)


def write_plain_text(text: str, output_path: str | Path) -> str:
    cleaned = normalize_extracted_text(text)
    if not cleaned:
        raise AppException(message="未提取到可写出的文本内容", code=4002, status_code=400)
    target = Path(output_path)
    target.write_text(cleaned, encoding="utf-8")
    return str(target)


def sections_to_text(sections: list[str]) -> str:
    cleaned = [normalize_extracted_text(section) for section in sections if normalize_extracted_text(section)]
    if not cleaned:
        raise AppException(message="未提取到可写出的文本内容", code=4002, status_code=400)
    return "\n\n".join(cleaned)


def tables_to_sections(tables: list[tuple[str, list[list[str]]]]) -> list[str]:
    if not tables:
        raise AppException(message="没有可提取的表格内容", code=4002, status_code=400)

    sections: list[str] = []
    for table_name, rows in tables:
        lines = [table_name]
        for row in rows:
            lines.append(" | ".join(cell for cell in row))
        sections.append("\n".join(lines))
    return sections


def build_html_from_sections(sections: list[str], output_path: str | Path, title: str) -> str:
    cleaned = [normalize_extracted_text(section) for section in sections if normalize_extracted_text(section)]
    if not cleaned:
        raise AppException(message="没有可写入 HTML 的文本内容", code=4002, status_code=400)

    article_parts: list[str] = []
    for index, section in enumerate(cleaned, start=1):
        if len(cleaned) > 1:
            article_parts.append(f"<h2>第 {index} 部分</h2>")
        for paragraph in section.split("\n"):
            if paragraph.strip():
                article_parts.append(f"<p>{escape(paragraph)}</p>")

    html = f"""<!doctype html>
<html lang="zh-CN">
  <head>
    <meta charset="utf-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1" />
    <title>{escape(title)}</title>
    <style>
      body {{
        margin: 0;
        background: #f4f7fb;
        color: #1f2937;
        font-family: "PingFang SC", "Microsoft YaHei", sans-serif;
      }}
      main {{
        max-width: 860px;
        margin: 40px auto;
        padding: 40px;
        background: #ffffff;
        border-radius: 24px;
        box-shadow: 0 18px 60px rgba(15, 23, 42, 0.08);
      }}
      h1 {{
        margin: 0 0 20px;
        font-size: 32px;
      }}
      h2 {{
        margin: 28px 0 14px;
        font-size: 22px;
      }}
      p {{
        margin: 0 0 12px;
        line-height: 1.85;
        white-space: pre-wrap;
        word-break: break-word;
      }}
    </style>
  </head>
  <body>
    <main>
      <h1>{escape(title)}</h1>
      {''.join(article_parts)}
    </main>
  </body>
</html>
"""
    target = Path(output_path)
    target.write_text(html, encoding="utf-8")
    return str(target)


def normalize_tabular_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def safe_filename(value: str, default: str) -> str:
    cleaned = re.sub(r"[^\w\u4e00-\u9fff.-]+", "-", value.strip(), flags=re.UNICODE).strip("-.")
    return cleaned or default


def extract_spreadsheet_sheets(input_path: str | Path) -> list[tuple[str, list[list[str]]]]:
    try:
        workbook = load_workbook(str(input_path), data_only=True)
    except Exception as exc:
        raise AppException(message=f"无法读取 Excel 文件：{exc}", code=4002, status_code=400) from exc

    sheets: list[tuple[str, list[list[str]]]] = []
    for worksheet in workbook.worksheets:
        rows: list[list[str]] = []
        for row in worksheet.iter_rows(values_only=True):
            values = [normalize_tabular_value(cell) for cell in row]
            if any(item != "" for item in values):
                rows.append(values)
        if rows:
            sheets.append((worksheet.title or f"Sheet{len(sheets) + 1}", rows))

    if not sheets:
        raise AppException(message="Excel 文件中没有可提取的数据", code=4002, status_code=400)
    return sheets


def write_csv_text(rows: list[list[str]], output_path: str | Path) -> str:
    target = Path(output_path)
    with target.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerows(rows)
    return str(target)


def build_csv_outputs(sheets: list[tuple[str, list[list[str]]]], output_dir: str | Path, basename: str) -> str:
    output_root = Path(output_dir)
    if len(sheets) == 1:
        sheet_name, rows = sheets[0]
        return write_csv_text(rows, output_root / f"{safe_filename(sheet_name, basename)}.csv")

    csv_dir = output_root / "csv-files"
    csv_dir.mkdir(parents=True, exist_ok=True)
    for index, (sheet_name, rows) in enumerate(sheets, start=1):
        write_csv_text(rows, csv_dir / f"{index:02d}-{safe_filename(sheet_name, 'sheet')}.csv")
    return zip_output_files(csv_dir, output_root / f"{basename}.zip")


def build_html_from_tables(tables: list[tuple[str, list[list[str]]]], output_path: str | Path, title: str) -> str:
    if not tables:
        raise AppException(message="没有可写入 HTML 的表格内容", code=4002, status_code=400)

    article_parts: list[str] = []
    for sheet_name, rows in tables:
        article_parts.append(f"<section><h2>{escape(sheet_name)}</h2><table>")
        for row_index, row in enumerate(rows):
            tag = "th" if row_index == 0 else "td"
            cells = "".join(f"<{tag}>{escape(cell)}</{tag}>" for cell in row)
            article_parts.append(f"<tr>{cells}</tr>")
        article_parts.append("</table></section>")

    html = f"""<!doctype html>
<html lang="zh-CN">
  <head>
    <meta charset="utf-8" />
    <meta name="viewport" content="width=device-width, initial-scale=1" />
    <title>{escape(title)}</title>
    <style>
      body {{
        margin: 0;
        background: #f4f7fb;
        color: #1f2937;
        font-family: "PingFang SC", "Microsoft YaHei", sans-serif;
      }}
      main {{
        max-width: 1080px;
        margin: 40px auto;
        padding: 40px;
        background: #ffffff;
        border-radius: 24px;
        box-shadow: 0 18px 60px rgba(15, 23, 42, 0.08);
      }}
      h1 {{
        margin: 0 0 24px;
        font-size: 32px;
      }}
      h2 {{
        margin: 28px 0 14px;
        font-size: 22px;
      }}
      table {{
        width: 100%;
        border-collapse: collapse;
        margin-bottom: 24px;
        overflow: hidden;
      }}
      th, td {{
        border: 1px solid #dbe4ee;
        padding: 10px 12px;
        text-align: left;
        vertical-align: top;
        word-break: break-word;
      }}
      th {{
        background: #edf4ff;
      }}
    </style>
  </head>
  <body>
    <main>
      <h1>{escape(title)}</h1>
      {''.join(article_parts)}
    </main>
  </body>
</html>
"""
    target = Path(output_path)
    target.write_text(html, encoding="utf-8")
    return str(target)


def extract_csv_rows(input_path: str | Path) -> list[list[str]]:
    text = read_text_content(input_path)
    reader = csv.reader(StringIO(text))
    rows = [list(row) for row in reader]
    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if not rows:
        raise AppException(message="CSV 文件中没有可提取的数据", code=4002, status_code=400)
    return rows


def build_excel_from_rows(rows: list[list[str]], output_path: str | Path, sheet_name: str = "Sheet1") -> str:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = safe_filename(sheet_name, "Sheet1")[:31]
    for row in rows:
        worksheet.append(row)
    target = Path(output_path)
    workbook.save(str(target))
    return str(target)


def extract_ppt_sections(input_path: str | Path) -> list[str]:
    try:
        presentation = Presentation(str(input_path))
    except Exception as exc:
        raise AppException(message=f"无法读取 PPT 文件：{exc}", code=4002, status_code=400) from exc

    sections: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            text = normalize_extracted_text(text)
            if text:
                texts.append(text)
        if texts:
            sections.append(f"第 {index} 页\n" + "\n".join(texts))

    if not sections:
        raise AppException(message="PPT 文件中没有可提取的文本内容", code=4002, status_code=400)
    return sections


def split_text_for_pdf(text: str, font_name: str, font_size: int, max_width: float) -> list[str]:
    if not text:
        return [""]

    rows: list[str] = []
    current = ""
    for char in text:
        candidate = f"{current}{char}"
        if pdfmetrics.stringWidth(candidate, font_name, font_size) <= max_width:
            current = candidate
            continue
        if current:
            rows.append(current)
            current = char
        else:
            rows.append(candidate)
            current = ""
    if current:
        rows.append(current)
    return rows or [""]


def write_text_pdf(sections: list[str], output_path: str | Path, title: str) -> str:
    cleaned = [normalize_extracted_text(section) for section in sections if normalize_extracted_text(section)]
    if not cleaned:
        raise AppException(message="未提取到可转换的文本内容", code=4002, status_code=400)

    target = Path(output_path)
    font_name = get_pdf_font_name()
    page_width, page_height = A4
    margin_x = 48
    margin_top = 52
    margin_bottom = 48
    title_size = 18
    body_size = 11
    line_height = body_size + 7
    max_width = page_width - margin_x * 2

    pdf = canvas.Canvas(str(target), pagesize=A4)
    pdf.setTitle(title)

    def start_page(with_title: bool) -> float:
        y = page_height - margin_top
        if with_title:
            pdf.setFont(font_name, title_size)
            pdf.drawString(margin_x, y, title)
            y -= title_size + 14
        pdf.setFont(font_name, body_size)
        return y

    y = start_page(with_title=True)
    for index, section in enumerate(cleaned, start=1):
        if len(cleaned) > 1:
            heading = f"第 {index} 部分"
            if y < margin_bottom + line_height * 2:
                pdf.showPage()
                y = start_page(with_title=False)
            pdf.setFont(font_name, 13)
            pdf.drawString(margin_x, y, heading)
            y -= 22
            pdf.setFont(font_name, body_size)

        paragraphs = section.split("\n")
        for paragraph in paragraphs:
            wrapped = split_text_for_pdf(paragraph or " ", font_name, body_size, max_width)
            for row in wrapped:
                if y < margin_bottom + line_height:
                    pdf.showPage()
                    y = start_page(with_title=False)
                pdf.drawString(margin_x, y, row)
                y -= line_height
            y -= 2
        y -= 8

    pdf.save()
    return str(target)


def extract_docx_sections(input_path: str | Path) -> list[str]:
    try:
        document = Document(str(input_path))
    except Exception as exc:
        raise AppException(message=f"无法读取 Word 文档：{exc}", code=4002, status_code=400) from exc

    sections: list[str] = []
    paragraphs = [item.text.strip() for item in document.paragraphs if item.text.strip()]
    if paragraphs:
        sections.append("\n".join(paragraphs))

    table_rows: list[str] = []
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                table_rows.append(" | ".join(values))
    if table_rows:
        sections.append("\n".join(table_rows))

    if not sections:
        raise AppException(message="Word 文档中没有可转换的文本内容", code=4002, status_code=400)
    return sections


def extract_pdf_sections(input_path: str | Path) -> list[str]:
    try:
        reader = PdfReader(str(input_path))
    except Exception as exc:
        raise AppException(message=f"无法读取 PDF 文件：{exc}", code=4002, status_code=400) from exc

    sections: list[str] = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        text = normalize_extracted_text(text)
        if text:
            sections.append(text)

    if not sections:
        raise AppException(message="PDF 中没有可提取的文本内容", code=4002, status_code=400)
    return sections


def build_docx_from_sections(sections: list[str], output_path: str | Path, title: str) -> str:
    cleaned = [normalize_extracted_text(section) for section in sections if normalize_extracted_text(section)]
    if not cleaned:
        raise AppException(message="没有可写入 Word 的文本内容", code=4002, status_code=400)

    document = Document()
    document.core_properties.title = title
    document.add_heading(title, level=0)

    for index, section in enumerate(cleaned, start=1):
        if len(cleaned) > 1:
            document.add_heading(f"第 {index} 部分", level=1)
        for paragraph in section.split("\n"):
            if paragraph.strip():
                document.add_paragraph(paragraph)
        if index != len(cleaned):
            document.add_page_break()

    target = Path(output_path)
    document.save(str(target))
    return str(target)


def render_pdf_pages_to_jpg(input_path: str | Path, output_dir: str | Path, scale: float = 2.0, quality: int = 90) -> str:
    target_dir = Path(output_dir) / "pdf-pages"
    target_dir.mkdir(parents=True, exist_ok=True)

    try:
        pdf = fitz.open(str(input_path))
    except Exception as exc:
        raise AppException(message=f"无法读取 PDF 文件：{exc}", code=4002, status_code=400) from exc

    if pdf.page_count == 0:
        raise AppException(message="PDF 文件没有可转换的页面", code=4002, status_code=400)

    matrix = fitz.Matrix(float(scale), float(scale))
    for index, page in enumerate(pdf, start=1):
        pixmap = page.get_pixmap(matrix=matrix, alpha=False)
        image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
        image.save(target_dir / f"page-{index}.jpg", format="JPEG", quality=max(60, min(int(quality), 95)), optimize=True)

    pdf.close()
    return zip_output_files(target_dir, Path(output_dir) / "pdf-pages.zip")


def image_to_pdf(input_path: str | Path, output_path: str | Path) -> str:
    try:
        image = Image.open(str(input_path))
        image.load()
    except Exception as exc:
        raise AppException(message=f"无法读取图片文件：{exc}", code=4002, status_code=400) from exc

    target = Path(output_path)
    ensure_rgb(image).save(str(target), format="PDF", resolution=150.0)
    return str(target)


def heic_to_pdf(input_path: str | Path, output_path: str | Path) -> str:
    register_heic_support()
    return image_to_pdf(input_path, output_path)


def extract_epub_sections(input_path: str | Path) -> tuple[str, list[str]]:
    try:
        book = epub.read_epub(str(input_path))
    except Exception as exc:
        raise AppException(message=f"无法读取 EPUB 文件：{exc}", code=4002, status_code=400) from exc

    metadata = book.get_metadata("DC", "title")
    title = metadata[0][0] if metadata and metadata[0] else Path(input_path).stem

    sections: list[str] = []
    for item in book.get_items_of_type(ITEM_DOCUMENT):
        soup = BeautifulSoup(item.get_content(), "html.parser")
        for node in soup(["script", "style"]):
            node.decompose()
        text = normalize_extracted_text(soup.get_text("\n"))
        if text:
            sections.append(text)

    if not sections:
        raise AppException(message="EPUB 中没有可提取的文本内容", code=4002, status_code=400)
    return title, sections


def build_epub_from_sections(sections: list[str], output_path: str | Path, title: str) -> str:
    cleaned = [normalize_extracted_text(section) for section in sections if normalize_extracted_text(section)]
    if not cleaned:
        raise AppException(message="没有可写入 EPUB 的文本内容", code=4002, status_code=400)

    book = epub.EpubBook()
    book.set_identifier(uuid4().hex)
    book.set_title(title)
    book.set_language("zh-CN")

    chapters = []
    for index, section in enumerate(cleaned, start=1):
        chapter = epub.EpubHtml(title=f"第 {index} 部分", file_name=f"chapter-{index}.xhtml", lang="zh-CN")
        paragraphs = [f"<p>{escape(line)}</p>" for line in section.split("\n") if line.strip()]
        chapter.content = f"<h1>第 {index} 部分</h1>{''.join(paragraphs) or '<p>空白内容</p>'}"
        book.add_item(chapter)
        chapters.append(chapter)

    book.toc = tuple(chapters)
    book.spine = ["nav", *chapters]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())

    target = Path(output_path)
    epub.write_epub(str(target), book)
    return str(target)
