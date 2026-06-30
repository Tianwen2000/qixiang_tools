"""文件说明：实现「XLSX 转 JSON」工具的后端逻辑。"""

import json
from pathlib import Path

from openpyxl import load_workbook

from app.core.exceptions import AppException


TOOL_META = {
    "slug": "xlsx-to-json",
    "name": "XLSX 转 JSON",
    "category": "dev",
    "input_mode": "file",
    "result_type": "text",
}


def _sheet_to_rows(sheet) -> list[dict]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    header_row = rows[0]
    headers = []
    for index, cell in enumerate(header_row, start=1):
        label = str(cell).strip() if cell is not None and str(cell).strip() else f"column_{index}"
        headers.append(label)

    result = []
    for row in rows[1:]:
        item = {}
        for index, value in enumerate(row):
            if value is None:
                continue
            item[headers[index]] = value
        if item:
            result.append(item)
    return result


def run(input_path: str, **_: dict) -> str:
    path = Path(input_path)
    try:
        workbook = load_workbook(path, data_only=True)
    except Exception as exc:
        raise AppException(message=f"无法读取 XLSX 文件: {exc}", code=4002, status_code=400) from exc

    payload = {sheet.title: _sheet_to_rows(sheet) for sheet in workbook.worksheets}
    if len(payload) == 1:
        return json.dumps(next(iter(payload.values())), ensure_ascii=False, indent=2)
    return json.dumps(payload, ensure_ascii=False, indent=2)
