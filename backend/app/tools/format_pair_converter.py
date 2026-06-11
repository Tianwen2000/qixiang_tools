import csv
from io import StringIO
from pathlib import Path
from typing import Callable

from bs4 import BeautifulSoup

from app.core.exceptions import AppException
from app.tools import (
    csv_to_excel,
    csv_to_epub,
    csv_to_html,
    csv_to_pdf,
    csv_to_txt,
    csv_to_word,
    docx_to_pdf,
    epub_to_pdf,
    epub_to_txt,
    epub_to_word,
    excel_to_csv,
    excel_to_epub,
    excel_to_html,
    excel_to_pdf,
    excel_to_txt,
    excel_to_word,
    jpg_to_pdf,
    pdf_to_epub,
    pdf_to_html,
    pdf_to_jpg,
    pdf_to_txt,
    pdf_to_word,
    ppt_to_epub,
    ppt_to_html,
    ppt_to_pdf,
    ppt_to_word,
    txt_to_epub,
    txt_to_html,
    txt_to_pdf,
    txt_to_word,
    word_to_epub,
    word_to_html,
    word_to_txt,
)
from app.utils.document_utils import (
    build_epub_from_sections,
    build_docx_from_sections,
    build_excel_from_rows,
    extract_docx_sections,
    extract_epub_sections,
    extract_pdf_sections,
    extract_ppt_sections,
    write_text_pdf,
    read_plain_text,
    read_text_content,
    sections_to_text,
    tables_to_sections,
    write_csv_text,
    write_plain_text,
)
from pptx import Presentation


Handler = Callable[..., str]


def dispatch_pair(input_path: str, output_dir: str, direction: str, handlers: dict[str, tuple[set[str], Handler]], **params: dict) -> str:
    if direction not in handlers:
        raise AppException(message="不支持的转换方向", code=4001, status_code=400)

    allowed_suffixes, handler = handlers[direction]
    suffix = Path(input_path).suffix.lower().lstrip(".")
    if suffix not in allowed_suffixes:
        labels = " / ".join(f".{item}" for item in sorted(allowed_suffixes))
        raise AppException(message=f"当前方向请上传 {labels} 文件", code=4002, status_code=400)

    return handler(input_path=input_path, output_dir=output_dir, **params)


def html_to_text(input_path: str, output_dir: str, **_: dict) -> str:
    soup = BeautifulSoup(read_text_content(input_path), "html.parser")
    text = soup.get_text("\n")
    return write_plain_text(text, Path(output_dir) / "html-to-txt.txt")


def html_to_word(input_path: str, output_dir: str, title: str = "", **_: dict) -> str:
    soup = BeautifulSoup(read_text_content(input_path), "html.parser")
    text = soup.get_text("\n")
    source_title = soup.title.string.strip() if soup.title and soup.title.string else ""
    doc_title = title.strip() or source_title or "HTML 转 Word"
    return build_docx_from_sections([text], Path(output_dir) / "html-to-word.docx", title=doc_title)


def html_to_pdf(input_path: str, output_dir: str, title: str = "", **_: dict) -> str:
    soup = BeautifulSoup(read_text_content(input_path), "html.parser")
    text = soup.get_text("\n")
    source_title = soup.title.string.strip() if soup.title and soup.title.string else ""
    pdf_title = title.strip() or source_title or "HTML 转 PDF"
    return write_text_pdf([text], Path(output_dir) / "html-to-pdf.pdf", title=pdf_title)


def _rows_from_sections(sections: list[str]) -> list[list[str]]:
    text = sections_to_text(sections)
    rows: list[list[str]] = []
    for line in text.splitlines():
        value = line.strip()
        if not value:
            continue
        if "|" in value:
            rows.append([cell.strip() for cell in value.split("|")])
        elif "\t" in value:
            rows.append([cell.strip() for cell in value.split("\t")])
        elif "," in value:
            rows.extend(csv.reader(StringIO(value)))
        else:
            rows.append([value])
    if not rows:
        raise AppException(message="没有可转换的数据", code=4002, status_code=400)
    return rows


def pdf_to_csv(input_path: str, output_dir: str, **_: dict) -> str:
    return write_csv_text(_rows_from_sections(extract_pdf_sections(input_path)), Path(output_dir) / "pdf-to-csv.csv")


def pdf_to_excel(input_path: str, output_dir: str, sheet_name: str = "", **_: dict) -> str:
    return build_excel_from_rows(_rows_from_sections(extract_pdf_sections(input_path)), Path(output_dir) / "pdf-to-excel.xlsx", sheet_name=sheet_name or "PDF 数据")


def word_to_csv(input_path: str, output_dir: str, **_: dict) -> str:
    return write_csv_text(_rows_from_sections(extract_docx_sections(input_path)), Path(output_dir) / "word-to-csv.csv")


def word_to_excel(input_path: str, output_dir: str, sheet_name: str = "", **_: dict) -> str:
    return build_excel_from_rows(_rows_from_sections(extract_docx_sections(input_path)), Path(output_dir) / "word-to-excel.xlsx", sheet_name=sheet_name or "Word 数据")


def epub_to_csv(input_path: str, output_dir: str, **_: dict) -> str:
    _, sections = extract_epub_sections(input_path)
    return write_csv_text(_rows_from_sections(sections), Path(output_dir) / "epub-to-csv.csv")


def epub_to_excel(input_path: str, output_dir: str, sheet_name: str = "", **_: dict) -> str:
    _, sections = extract_epub_sections(input_path)
    return build_excel_from_rows(_rows_from_sections(sections), Path(output_dir) / "epub-to-excel.xlsx", sheet_name=sheet_name or "EPUB 数据")


def _rows_from_plain_text(input_path: str) -> list[list[str]]:
    text = read_plain_text(input_path)
    rows: list[list[str]] = []
    for line in text.splitlines():
        value = line.strip()
        if not value:
            continue
        if "\t" in value:
            rows.append([cell.strip() for cell in value.split("\t")])
        elif "," in value:
            rows.extend(csv.reader(StringIO(value)))
        else:
            rows.append([value])
    if not rows:
        raise AppException(message="TXT 文件中没有可转换的数据", code=4002, status_code=400)
    return rows


def txt_to_csv(input_path: str, output_dir: str, **_: dict) -> str:
    return write_csv_text(_rows_from_plain_text(input_path), Path(output_dir) / "txt-to-csv.csv")


def txt_to_excel(input_path: str, output_dir: str, sheet_name: str = "", **_: dict) -> str:
    return build_excel_from_rows(_rows_from_plain_text(input_path), Path(output_dir) / "txt-to-excel.xlsx", sheet_name=sheet_name or "TXT 数据")


def _html_tables(input_path: str) -> list[tuple[str, list[list[str]]]]:
    soup = BeautifulSoup(read_text_content(input_path), "html.parser")
    tables: list[tuple[str, list[list[str]]]] = []
    for index, table in enumerate(soup.find_all("table"), start=1):
        rows: list[list[str]] = []
        for tr in table.find_all("tr"):
            cells = [cell.get_text(" ", strip=True) for cell in tr.find_all(["th", "td"])]
            if any(cells):
                rows.append(cells)
        if rows:
            tables.append((f"表格{index}", rows))

    if tables:
        return tables

    text_rows = [[line] for line in soup.get_text("\n").splitlines() if line.strip()]
    if not text_rows:
        raise AppException(message="HTML 文件中没有可转换的数据", code=4002, status_code=400)
    return [("HTML 文本", text_rows)]


def html_to_csv(input_path: str, output_dir: str, **_: dict) -> str:
    tables = _html_tables(input_path)
    if len(tables) == 1:
        return write_csv_text(tables[0][1], Path(output_dir) / "html-to-csv.csv")

    csv_dir = Path(output_dir) / "html-csv-files"
    csv_dir.mkdir(parents=True, exist_ok=True)
    from app.utils.image_utils import zip_output_files

    for index, (_, rows) in enumerate(tables, start=1):
        write_csv_text(rows, csv_dir / f"table-{index:02d}.csv")
    return zip_output_files(csv_dir, Path(output_dir) / "html-to-csv.zip")


def html_to_excel(input_path: str, output_dir: str, sheet_name: str = "", **_: dict) -> str:
    tables = _html_tables(input_path)
    rows: list[list[str]] = []
    for title, table_rows in tables:
        if rows:
            rows.append([])
        rows.append([title])
        rows.extend(table_rows)
    return build_excel_from_rows(rows, Path(output_dir) / "html-to-excel.xlsx", sheet_name=sheet_name or "HTML 数据")


def _build_ppt_from_sections(sections: list[str], output_path: Path, title: str) -> str:
    cleaned = [section.strip() for section in sections if section.strip()]
    if not cleaned:
        raise AppException(message="没有可写入 PPT 的文本内容", code=4002, status_code=400)

    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = title
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = "由格式转换工具生成"

    for index, section in enumerate(cleaned, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"第 {index} 部分"
        body = slide.placeholders[1]
        body.text = section[:2500]

    presentation.save(output_path)
    return str(output_path)


def txt_to_ppt(input_path: str, output_dir: str, title: str = "", **_: dict) -> str:
    text = read_plain_text(input_path)
    return _build_ppt_from_sections([text], Path(output_dir) / "txt-to-ppt.pptx", title.strip() or "TXT 转 PPT")


def html_to_ppt(input_path: str, output_dir: str, title: str = "", **_: dict) -> str:
    soup = BeautifulSoup(read_text_content(input_path), "html.parser")
    source_title = soup.title.string.strip() if soup.title and soup.title.string else ""
    return _build_ppt_from_sections([soup.get_text("\n")], Path(output_dir) / "html-to-ppt.pptx", title.strip() or source_title or "HTML 转 PPT")


def word_to_ppt(input_path: str, output_dir: str, title: str = "", **_: dict) -> str:
    return _build_ppt_from_sections(extract_docx_sections(input_path), Path(output_dir) / "word-to-ppt.pptx", title.strip() or "Word 转 PPT")


def pdf_to_ppt(input_path: str, output_dir: str, title: str = "", **_: dict) -> str:
    return _build_ppt_from_sections(extract_pdf_sections(input_path), Path(output_dir) / "pdf-to-ppt.pptx", title.strip() or "PDF 转 PPT")


def epub_to_ppt(input_path: str, output_dir: str, title: str = "", **_: dict) -> str:
    epub_title, sections = extract_epub_sections(input_path)
    return _build_ppt_from_sections(sections, Path(output_dir) / "epub-to-ppt.pptx", title.strip() or epub_title or "EPUB 转 PPT")


WORD_PDF_HANDLERS = {
    "word_to_pdf": ({"docx"}, docx_to_pdf.run),
    "pdf_to_word": ({"pdf"}, pdf_to_word.run),
}
PDF_JPG_HANDLERS = {
    "pdf_to_jpg": ({"pdf"}, pdf_to_jpg.run),
    "jpg_to_pdf": ({"jpg", "jpeg", "png", "webp", "bmp"}, jpg_to_pdf.run),
}
PDF_EPUB_HANDLERS = {
    "pdf_to_epub": ({"pdf"}, pdf_to_epub.run),
    "epub_to_pdf": ({"epub"}, epub_to_pdf.run),
}
PDF_TXT_HANDLERS = {
    "pdf_to_txt": ({"pdf"}, pdf_to_txt.run),
    "txt_to_pdf": ({"txt"}, txt_to_pdf.run),
}
PDF_HTML_HANDLERS = {
    "pdf_to_html": ({"pdf"}, pdf_to_html.run),
    "html_to_pdf": ({"html", "htm"}, html_to_pdf),
}
WORD_TXT_HANDLERS = {
    "word_to_txt": ({"docx"}, word_to_txt.run),
    "txt_to_word": ({"txt"}, txt_to_word.run),
}
TXT_EPUB_HANDLERS = {
    "txt_to_epub": ({"txt"}, txt_to_epub.run),
    "epub_to_txt": ({"epub"}, epub_to_txt.run),
}
WORD_EPUB_HANDLERS = {
    "word_to_epub": ({"docx"}, word_to_epub.run),
    "epub_to_word": ({"epub"}, epub_to_word.run),
}
WORD_HTML_HANDLERS = {
    "word_to_html": ({"docx"}, word_to_html.run),
    "html_to_word": ({"html", "htm"}, html_to_word),
}
TXT_HTML_HANDLERS = {
    "txt_to_html": ({"txt"}, txt_to_html.run),
    "html_to_txt": ({"html", "htm"}, html_to_text),
}
EXCEL_CSV_HANDLERS = {
    "excel_to_csv": ({"xlsx"}, excel_to_csv.run),
    "csv_to_excel": ({"csv"}, csv_to_excel.run),
}
EXCEL_TXT_HANDLERS = {
    "excel_to_txt": ({"xlsx"}, excel_to_txt.run),
    "txt_to_excel": ({"txt"}, txt_to_excel),
}
CSV_TXT_HANDLERS = {
    "csv_to_txt": ({"csv"}, csv_to_txt.run),
    "txt_to_csv": ({"txt"}, txt_to_csv),
}
EXCEL_HTML_HANDLERS = {
    "excel_to_html": ({"xlsx"}, excel_to_html.run),
    "html_to_excel": ({"html", "htm"}, html_to_excel),
}
CSV_HTML_HANDLERS = {
    "csv_to_html": ({"csv"}, csv_to_html.run),
    "html_to_csv": ({"html", "htm"}, html_to_csv),
}
EXCEL_PDF_HANDLERS = {
    "excel_to_pdf": ({"xlsx"}, excel_to_pdf.run),
    "pdf_to_excel": ({"pdf"}, pdf_to_excel),
}
CSV_PDF_HANDLERS = {
    "csv_to_pdf": ({"csv"}, csv_to_pdf.run),
    "pdf_to_csv": ({"pdf"}, pdf_to_csv),
}
EXCEL_WORD_HANDLERS = {
    "excel_to_word": ({"xlsx"}, excel_to_word.run),
    "word_to_excel": ({"docx"}, word_to_excel),
}
CSV_WORD_HANDLERS = {
    "csv_to_word": ({"csv"}, csv_to_word.run),
    "word_to_csv": ({"docx"}, word_to_csv),
}
EXCEL_EPUB_HANDLERS = {
    "excel_to_epub": ({"xlsx"}, excel_to_epub.run),
    "epub_to_excel": ({"epub"}, epub_to_excel),
}
CSV_EPUB_HANDLERS = {
    "csv_to_epub": ({"csv"}, csv_to_epub.run),
    "epub_to_csv": ({"epub"}, epub_to_csv),
}
PPT_TXT_HANDLERS = {
    "ppt_to_txt": ({"pptx"}, lambda input_path, output_dir, **params: write_plain_text(sections_to_text(extract_ppt_sections(input_path)), Path(output_dir) / "ppt-to-txt.txt")),
    "txt_to_ppt": ({"txt"}, txt_to_ppt),
}
PPT_HTML_HANDLERS = {
    "ppt_to_html": ({"pptx"}, ppt_to_html.run),
    "html_to_ppt": ({"html", "htm"}, html_to_ppt),
}
PPT_WORD_HANDLERS = {
    "ppt_to_word": ({"pptx"}, ppt_to_word.run),
    "word_to_ppt": ({"docx"}, word_to_ppt),
}
PPT_PDF_HANDLERS = {
    "ppt_to_pdf": ({"pptx"}, ppt_to_pdf.run),
    "pdf_to_ppt": ({"pdf"}, pdf_to_ppt),
}
PPT_EPUB_HANDLERS = {
    "ppt_to_epub": ({"pptx"}, ppt_to_epub.run),
    "epub_to_ppt": ({"epub"}, epub_to_ppt),
}
