import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.core.exceptions import AppException


def _normalize_rows(value: object) -> list[dict]:
    if isinstance(value, dict):
        return [value]
    if isinstance(value, list):
        rows = []
        for item in value:
            if isinstance(item, dict):
                rows.append(item)
            else:
                rows.append({"value": item})
        return rows
    raise AppException(message="JSON 内容必须是对象或数组", code=4001, status_code=400)


def _json_to_xlsx(input_path: str, output_dir: str, sheet_name: str = "") -> str:
    try:
        value = json.loads(Path(input_path).read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise AppException(message=f"JSON 格式无效：{exc.msg}", code=4001, status_code=400) from exc

    rows = _normalize_rows(value)
    headers = sorted({key for row in rows for key in row.keys()}) or ["value"]

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = (sheet_name or "Sheet1").strip()[:31] or "Sheet1"
    worksheet.append(headers)
    for row in rows:
        worksheet.append([json.dumps(row.get(header, ""), ensure_ascii=False) if isinstance(row.get(header, ""), (dict, list)) else row.get(header, "") for header in headers])

    target_path = Path(output_dir) / "json-to-xlsx.xlsx"
    workbook.save(target_path)
    return str(target_path)


def _sheet_to_rows(sheet) -> list[dict]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(cell).strip() if cell is not None and str(cell).strip() else f"column_{index}" for index, cell in enumerate(rows[0], start=1)]
    result = []
    for row in rows[1:]:
        item = {}
        for index, value in enumerate(row):
            if value is not None:
                item[headers[index]] = value
        if item:
            result.append(item)
    return result


def _xlsx_to_json(input_path: str, output_dir: str) -> str:
    try:
        workbook = load_workbook(input_path, data_only=True)
    except Exception as exc:
        raise AppException(message=f"无法读取 XLSX 文件: {exc}", code=4002, status_code=400) from exc

    payload = {sheet.title: _sheet_to_rows(sheet) for sheet in workbook.worksheets}
    if len(payload) == 1:
        payload = next(iter(payload.values()))

    target_path = Path(output_dir) / "xlsx-to-json.json"
    target_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return str(target_path)


def run(input_path: str, output_dir: str, direction: str = "json_to_xlsx", sheet_name: str = "", **_: dict) -> str:
    suffix = Path(input_path).suffix.lower().lstrip(".")
    if direction == "json_to_xlsx":
        if suffix != "json":
            raise AppException(message="当前方向请上传 .json 文件", code=4002, status_code=400)
        return _json_to_xlsx(input_path, output_dir, sheet_name=sheet_name)
    if direction == "xlsx_to_json":
        if suffix != "xlsx":
            raise AppException(message="当前方向请上传 .xlsx 文件", code=4002, status_code=400)
        return _xlsx_to_json(input_path, output_dir)
    raise AppException(message="不支持的转换方向", code=4001, status_code=400)
